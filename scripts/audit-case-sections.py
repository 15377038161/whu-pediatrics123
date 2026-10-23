"""Compare section repairs against reproducible local cases; stage candidates only.

Uses the pre-fix generator at the recorded Git baseline to identify the exact
primary document. Never overwrites knowledge/, imports to DB, or publishes.
"""
import argparse
import ast
from collections import Counter
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
import re
import subprocess

from case_source_sections import split_case_sections

ROOT = Path(__file__).resolve().parents[1]


def digest(path):
    return sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', default='outputs/case-library-review/20261021')
    args = parser.parse_args()
    out = (ROOT / args.output).resolve()
    if ROOT not in out.parents or out == ROOT:
        raise SystemExit('Output must stay in repository.')
    if subprocess.run(['git', 'check-ignore', str(out.relative_to(ROOT))], cwd=ROOT, capture_output=True).returncode:
        raise SystemExit('Output must be ignored.')
    inventory = json.loads((out / 'scope.json').read_text(encoding='utf-8'))
    baseline = inventory['baseline']['gitHead']
    if subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip() != baseline:
        raise SystemExit('Refresh inventory after synchronizing Git before auditing sections.')
    material_index = {}
    for line in (out / 'materials.jsonl').read_text(encoding='utf-8').splitlines():
        material = json.loads(line)
        if material['role'] == 'source_material':
            material_index.setdefault(material['sha256'], []).append(material)
    old_generators = []
    source_cleaner = None
    for source_path in ('scripts/prepare-teacher-cases.py', 'scripts/build-runtime-cases-from-handover.py'):
        original = subprocess.check_output(['git', 'show', f'{baseline}:{source_path}'], cwd=ROOT, encoding='utf-8')
        function = next(node for node in ast.parse(original).body if isinstance(node, ast.FunctionDef) and node.name == 'sections')
        namespace = {'re': re}
        exec(compile(ast.Module(body=[function], type_ignores=[]), source_path, 'exec'), namespace)
        old_generators.append((source_path, namespace['sections']))
        if source_path.endswith('prepare-teacher-cases.py'):
            cleaner = next(node for node in ast.parse(original).body if isinstance(node, ast.FunctionDef) and node.name == 'clean')
            exec(compile(ast.Module(body=[cleaner], type_ignores=[]), source_path, 'exec'), namespace)
            source_cleaner = namespace['clean']
    candidate_dir = out / 'section-candidates'
    candidate_dir.mkdir(parents=True, exist_ok=True)
    results = []
    for file in sorted((ROOT / 'knowledge/teaching-cases').glob('*.json')):
        source_sha = digest(file)
        record = json.loads(file.read_text(encoding='utf-8'))
        current = record['sections']
        matches = []
        for document in record.get('documents', []):
            for generator, split_old in old_generators:
                if split_old(document['text']) == current:
                    matches.append((generator, document))
        result = {'caseId': record['caseId'], 'baselineRuntimeSha256': source_sha,
                  'oldSectionLengths': {key: len(value) for key, value in current.items()},
                  'status': 'not_reproduced_from_recorded_primary_document', 'applied': False}
        if matches:
            generator, document = matches[0]
            new = split_case_sections(document['text'])
            if generator.endswith('prepare-teacher-cases.py'):
                new['history'] = re.sub(r'(?:门诊以|门诊拟|考虑诊断为|外院诊断为).*(?:收入院|入院|住院)[。；;]?', '', new['history']).strip()
            result.update(generator=generator, primaryDocumentSha256=document['sha256'],
                          newSectionLengths={key: len(value) for key, value in new.items()},
                          changedSections=[key for key in current if current[key] != new[key]])
            source_matches = [m for m in material_index.get(document['sha256'], []) if m['sourceCaseId'] == record['caseId']]
            result['verifiedOriginalMaterialIds'] = [m['materialId'] for m in source_matches]
            result['originalTextComparison'] = 'not_verified'
            if source_matches:
                source_file = ROOT / source_matches[0]['path']
                if source_file.suffix.lower() == '.docx':
                    from docx import Document
                    original_doc = Document(source_file)
                    lines = [p.text.strip() for p in original_doc.paragraphs if p.text.strip()]
                    for table in original_doc.tables:
                        for row in table.rows:
                            line = ' | '.join(cell.text.strip().replace('\n', ' / ') for cell in row.cells)
                            if line.strip(' |'):
                                lines.append(line)
                    extracted = source_cleaner('\n'.join(lines))
                    result['originalTextComparison'] = 'exact_cleaned_text_match' if extracted == document['text'] else 'text_mismatch'
            # Compare the multisets of complete source lines, not character counts.
            before = Counter(line for value in current.values() for line in value.splitlines() if line.strip())
            after = Counter(line for value in new.values() for line in value.splitlines() if line.strip())
            result['sourceLinesPreserved'] = before == after
            if new == current:
                result['status'] = 'unchanged'
            elif before != after:
                result['status'] = 'blocked_source_line_change'
            elif not source_matches:
                result['status'] = 'blocked_original_fingerprint_unverified'
            elif result['originalTextComparison'] == 'text_mismatch':
                result['status'] = 'blocked_original_text_mismatch'
            else:
                candidate = deepcopy(record)
                candidate['sections'] = new
                target = candidate_dir / file.name
                serialized = json.dumps(candidate, ensure_ascii=False, indent=2) + '\n'
                # Hash exactly the emitted bytes on Windows too (no CRLF translation).
                target.write_bytes(serialized.encode('utf-8'))
                result.update(status='candidate_staged_not_applied', candidate=target.relative_to(ROOT).as_posix(),
                              candidateSha256=sha256(serialized.encode('utf-8')).hexdigest(),
                              versionRisk='document_hash_unchanged_requires_runtime_content_version_fix')
        if digest(file) != source_sha:
            raise SystemExit('Source changed during comparison; rerun safely.')
        results.append(result)
    counts = Counter(row['status'] for row in results)
    report = {'baseline': baseline, 'counts': dict(counts), 'sourceMutation': False, 'results': results}
    (out / 'section-audit.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    lines = ['# 病例分段源文核对', '', f'代码基线：`{baseline}`。', '',
             '仅对能够由旧生成器逐字段复现、原件指纹已匹配的病例生成修订候选；未覆盖运行文件、未导入或发布。', '',
             '| 病例 | 状态 | 旧处置字符数 | 候选处置字符数 | 原文行保留 |', '|---|---|---:|---:|---|']
    for row in results:
        if row['status'] != 'unchanged':
            lines.append(f'| {row["caseId"]} | {row["status"]} | {row["oldSectionLengths"]["plan"]} | {row.get("newSectionLengths", {}).get("plan", "—")} | {row.get("sourceLinesPreserved", "—")} |')
    (out / 'section-audit.md').write_text('\n'.join(lines) + '\n', encoding='utf-8')
    print(json.dumps({'counts': dict(counts), 'sourceMutation': False}, ensure_ascii=False))
    return 1 if any(key.startswith('blocked_') for key in counts) else 0


if __name__ == '__main__':
    raise SystemExit(main())
