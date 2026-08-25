from pathlib import Path
from datetime import date

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw

import build_report as br


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "教师汇报材料"
ASSETS = DOCS / "assets"
DOCS.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)

FEATURE_OUT = DOCS / "珞珈儿科智训_功能介绍_教师版.docx"
FLOW_OUT = DOCS / "珞珈儿科智训_整体业务流程_教师版.docx"


def add_arrow(draw, start, end, color="#42B9C8", width=8):
    draw.line([start, end], fill=color, width=width)
    x, y = end
    draw.polygon([(x, y), (x - 20, y - 13), (x - 20, y + 13)], fill=color)


def feature_map():
    img = Image.new("RGB", (1800, 1040), "#F3F7F9")
    d = ImageDraw.Draw(img)
    title, heading, body, small = br.font(46, True), br.font(28, True), br.font(22), br.font(19)
    d.text((70, 45), "珞珈儿科智训功能全景", font=title, fill="#0A3556")
    d.text((70, 108), "以儿科临床胜任力为主线，连接训练、评价与教师内容治理。", font=body, fill="#60727E")

    br.rounded(d, (570, 185, 1230, 330), "#0A3556")
    d.text((680, 215), "儿童全周期临床胜任力", font=heading, fill="white")
    d.text((720, 265), "问诊 · 查体 · 推理 · 沟通 · 安全", font=body, fill="#CDE9F2")

    cards = [
        (80, 410, "01", "问诊与可视化查体", "患儿/家长双角色\n身体部位点击与器材校验"),
        (505, 410, "02", "儿童传染病专项", "年龄与季节特点\n隔离、报告及院感规则"),
        (930, 410, "03", "儿科沟通训练", "患儿—家长—医生三方\n人文与风险沟通"),
        (1355, 410, "04", "OSCE标准化评价", "行为证据与关键错误\n评分、复盘与补练"),
    ]
    for x, y, no, name, desc in cards:
        br.rounded(d, (x, y, x + 365, y + 230), "#FFFFFF", "#D9E4E9", 3)
        d.text((x + 24, y + 22), no, font=body, fill="#42B9C8")
        d.text((x + 24, y + 72), name, font=heading, fill="#0A3556")
        d.multiline_text((x + 24, y + 130), desc, font=small, fill="#60727E", spacing=10)
        add_arrow(d, (900, 330), (x + 182, y - 12), width=5)

    lower = [
        (190, "教师内容管理", "病例、指南、量表、素材审核与版本管理"),
        (690, "动态病例与安全规则", "学生操作推动病情演变；危险动作受规则拦截"),
        (1190, "学习数据与持续改进", "全过程留痕、薄弱项统计、教学内容迭代"),
    ]
    for x, name, desc in lower:
        br.rounded(d, (x, 755, x + 420, 930), "#EAF4F8", "#9FC8DA", 3)
        d.text((x + 24, 790), name, font=heading, fill="#176FA9")
        d.multiline_text((x + 24, 845), desc, font=small, fill="#1B2E3B", spacing=8)
    path = ASSETS / "功能全景图.png"
    img.save(path)
    return path


def role_journey():
    img = Image.new("RGB", (1800, 1120), "white")
    d = ImageDraw.Draw(img)
    title, heading, body, small = br.font(44, True), br.font(26, True), br.font(21), br.font(18)
    d.text((65, 42), "三类角色协同路径", font=title, fill="#0A3556")
    stages = ["内容准备", "发布训练", "学生训练", "评价复盘", "持续改进"]
    xs = [260, 550, 840, 1130, 1420]
    for x, text in zip(xs, stages):
        br.rounded(d, (x, 135, x + 235, 205), "#0A3556")
        d.text((x + 45, 155), text, font=body, fill="white")
    lanes = [
        (270, "课程教师", "提供病例与量表", "配置任务和难度", "查看班级表现", "调整教学重点", "确认下一版本"),
        (545, "学生", "查看训练要求", "进入指定病例", "问诊/查体/决策", "查看证据反馈", "完成个性化补练"),
        (820, "审核教师", "审核医学内容", "锁定正式版本", "抽查异常记录", "复核争议评分", "批准内容更新"),
    ]
    colors = ["#EAF4F8", "#EAF7F4", "#FFF7E7"]
    for idx, (y, role, *items) in enumerate(lanes):
        br.rounded(d, (55, y, 240, y + 190), colors[idx], "#D9E4E9", 3)
        d.text((90, y + 70), role, font=heading, fill="#0A3556")
        for col, item in enumerate(items):
            x = xs[col]
            br.rounded(d, (x, y, x + 235, y + 190), colors[idx], "#D9E4E9", 2)
            d.multiline_text((x + 22, y + 58), item, font=small, fill="#1B2E3B", spacing=8, align="center")
            if col < 4:
                add_arrow(d, (x + 238, y + 95), (xs[col + 1] - 8, y + 95), width=5)
    d.text((65, 1040), "原则：教师决定内容和标准；学生完成训练；审核教师保证医学正确性与版本可追溯。", font=body, fill="#60727E")
    path = ASSETS / "角色协同路径.png"
    img.save(path)
    return path


def business_flow():
    img = Image.new("RGB", (1800, 1040), "#F3F7F9")
    d = ImageDraw.Draw(img)
    title, heading, body, small = br.font(44, True), br.font(26, True), br.font(20), br.font(17)
    d.text((65, 42), "整体业务闭环", font=title, fill="#0A3556")
    d.text((65, 102), "从老师提供资源，到学生训练，再回到课程改进。", font=body, fill="#60727E")
    steps = [
        ("01", "材料提供", "病例、指南、量表\n体征与沟通素材"),
        ("02", "知识加工", "脱敏、结构化\n标注、规则映射"),
        ("03", "医学审核", "来源核验、双人审核\n版本锁定"),
        ("04", "任务发布", "选择病例、难度\n时间与提示策略"),
        ("05", "学生训练", "问诊、可视化查体\n决策与沟通"),
        ("06", "OSCE评价", "行为证据、扣分原因\n补练建议"),
        ("07", "教学改进", "薄弱项分析\n病例和课程迭代"),
    ]
    coords = [(70,240),(500,240),(930,240),(1360,240),(1360,625),(930,625),(500,625)]
    for i, ((no, name, desc), (x,y)) in enumerate(zip(steps, coords)):
        fill = "#FFFFFF" if i < 4 else ("#EAF7F4" if i < 6 else "#FFF7E7")
        br.rounded(d, (x,y,x+340,y+220), fill, "#BFD5DF", 3)
        d.text((x+24,y+22),no,font=body,fill="#42B9C8")
        d.text((x+24,y+65),name,font=heading,fill="#0A3556")
        d.multiline_text((x+24,y+125),desc,font=small,fill="#60727E",spacing=8)
    for a,b in zip(coords[:3],coords[1:4]): add_arrow(d,(a[0]+345,a[1]+110),(b[0]-8,b[1]+110),width=6)
    add_arrow(d,(1530,465),(1530,615),width=6)
    add_arrow(d,(1352,735),(1278,735),width=6)
    add_arrow(d,(922,735),(848,735),width=6)
    add_arrow(d,(500,845),(240,845),width=6)
    br.rounded(d,(70,750,410,940),"#EAF4F8","#BFD5DF",3)
    d.text((95,790),"质量与安全贯穿",font=heading,fill="#176FA9")
    d.multiline_text((95,845),"隐私脱敏 · 来源追溯\n教师审核 · 权限与审计",font=small,fill="#1B2E3B",spacing=8)
    path = ASSETS / "整体业务闭环.png"
    img.save(path)
    return path


def exam_flow():
    img = Image.new("RGB", (1800, 970), "white")
    d = ImageDraw.Draw(img)
    title, heading, body, small = br.font(42, True), br.font(25, True), br.font(20), br.font(17)
    d.text((65, 42), "学生一次训练的完整路径", font=title, fill="#0A3556")
    steps = [
        ("进入病例", "阅读任务\n确认患儿信息"),
        ("双角色问诊", "询问患儿\n向家长补充病史"),
        ("观察与查体", "点击身体部位\n观察表面体征"),
        ("器材检查", "选器材＋匹配部位\n解锁深层信息"),
        ("临床决策", "诊断、检查\n处置与病情复评"),
        ("沟通宣教", "回应家长焦虑\n说明风险与计划"),
        ("评价补练", "OSCE证据报告\n进入薄弱项训练"),
    ]
    coords=[]
    for i in range(4): coords.append((65+i*430,200))
    for i in range(3): coords.append((1355-i*430,600))
    for idx, ((name,desc),(x,y)) in enumerate(zip(steps,coords)):
        br.rounded(d,(x,y,x+340,y+210),"#EAF4F8" if idx<4 else "#FFF7E7","#BFD5DF",3)
        d.text((x+25,y+28),f"{idx+1:02d}",font=body,fill="#42B9C8")
        d.text((x+25,y+70),name,font=heading,fill="#0A3556")
        d.multiline_text((x+25,y+125),desc,font=small,fill="#60727E",spacing=7)
    for a,b in zip(coords[:3],coords[1:4]): add_arrow(d,(a[0]+345,a[1]+105),(b[0]-8,b[1]+105),width=6)
    add_arrow(d,(1525,415),(1525,590),width=6)
    add_arrow(d,(1347,705),(1278,705),width=6)
    add_arrow(d,(917,705),(848,705),width=6)
    d.text((65,885),"任何步骤均记录操作顺序、用时、结果和错误类型，作为OSCE评价证据。",font=body,fill="#60727E")
    path = ASSETS / "学生训练路径.png"
    img.save(path)
    return path


def setup_doc(doc, short_title):
    br.setup_styles(doc)
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    br.set_run_font(hp.add_run(f"珞珈儿科智训  |  {short_title}"), size=8.5, color=br.MUTED)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    br.set_run_font(fp.add_run("武汉大学儿科学教学智能体  ·  "), size=8, color=br.MUTED)
    br.add_field(fp, "PAGE")


def add_cover(doc, title, subtitle, conclusion):
    br.add_para(doc,"武汉大学儿科学 · 智能体大赛项目",size=11,color=br.GOLD,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=36)
    br.add_para(doc,title,size=29,color=br.NAVY,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=12)
    br.add_para(doc,subtitle,size=15,color=br.BLUE,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=34)
    br.add_para(doc,"面向课程教师、医学审核教师与项目建设人员",size=11,color=br.MUTED,italic=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=56)
    br.add_callout(doc,"核心结论",conclusion,fill=br.PALE_BLUE)
    br.add_para(doc,f"编制日期：{date.today().strftime('%Y年%m月%d日')}",size=10,color=br.MUTED,align=WD_ALIGN_PARAGRAPH.CENTER,before=70,after=5)
    br.add_para(doc,"文档状态：教师沟通与建设确认稿",size=10,color=br.MUTED,align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def add_picture(doc, path, title, desc, caption):
    doc.add_picture(str(path), width=Inches(6.45))
    br.set_last_image_alt(doc, title, desc)
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    br.add_para(doc, caption, size=8.5, color=br.MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)


def remove_trailing_empty_paragraph(doc):
    if doc.paragraphs and not doc.paragraphs[-1].text.strip():
        paragraph = doc.paragraphs[-1]._element
        paragraph.getparent().remove(paragraph)


def build_feature_doc():
    fmap = feature_map()
    roles = role_journey()
    doc = Document()
    setup_doc(doc, "功能介绍（教师版）")
    add_cover(doc,"珞珈儿科智训：功能介绍","儿童全周期临床胜任力训练与OSCE考核智能体","系统不是普通医学问答工具，而是一套由教师控制内容与标准、由学生完成临床操作、由系统记录证据并形成反馈的儿科临床教学平台。")

    br.add_heading(doc,"一、系统定位与建设目标",1)
    br.add_para(doc,"“珞珈儿科智训”面向儿科学专业学生的见习、实习和OSCE训练。系统以虚拟患儿病例为载体，同时模拟患儿与家长，支持学生完成问诊、观察、查体、检查、诊断、处置和沟通，并依据教师确认的量表形成可解释评价。")
    br.add_callout(doc,"老师最需要了解的三点","第一，医学内容和评分标准由教师决定；第二，学生必须主动操作才能获得信息；第三，系统记录全过程，而不是只给一个总分。",fill=br.PALE_GOLD,accent=br.GOLD)
    br.add_heading(doc,"建设目标",2)
    for t in ["让难以重复配合的真实患儿场景能够安全、稳定地反复训练。","把患儿、家长和医生三方沟通纳入儿科临床能力评价。","让低频、高风险病例能够在无患者伤害的环境中练习。","让OSCE评分能够回到学生具体操作证据，并自动推荐补练。","沉淀武汉大学儿科学课程自己的病例、量表、指南和教学经验。"]: br.add_bullet(doc,t)

    br.add_heading(doc,"二、系统功能全景",1)
    add_picture(doc,fmap,"珞珈儿科智训功能全景图","四个核心功能模块以及教师内容管理、动态病例安全规则和学习数据改进能力。","图1  系统功能全景")
    br.add_heading(doc,"2.1 儿科问诊与可视化查体",2)
    br.add_para(doc,"系统分别模拟患儿和家长。患儿根据年龄表达主观感受，家长补充出生史、喂养史、发育史、接种史和病程时间线，避免两个角色重复提供全部信息。")
    for t in ["可视化患儿模型：点击口腔、胸部、腹部、皮肤等部位，触发张口、吐舌、转身、抬臂等动作或表面体征。","器材工具栏：听诊器、血压计、体温计、压舌板等器材由学生自主选择。","专业校验：只有“病例状态＋器材＋身体部位＋操作顺序”匹配，才开放呼吸音、血压、体温等深层信息。","操作反馈：提供选中状态、动作动画、成功提示和分级纠错，并记录错误次数、顺序和用时。"]: br.add_bullet(doc,t)
    br.add_heading(doc,"2.2 儿童传染病专项",2)
    for t in ["按年龄、季节、传播途径、临床综合征和隔离要求组织知识。","突出儿童非典型表现、疫苗接种、快速恶化风险和剂量安全。","训练接触史、聚集性发病史、报告流程、院感防控和家庭宣教。","每条知识记录来源、版本、审核教师和下次复核日期。"]: br.add_bullet(doc,t)
    br.add_heading(doc,"2.3 儿科沟通专项",2)
    for t in ["覆盖哭闹拒检、焦虑家长、抗菌药诉求、疫苗犹豫、知情同意和青春期隐私等场景。","评价共情、通俗表达、风险解释、共同决策、儿童权益保护和安全网告知。","同一病例可配置不同家长性格与认知水平，训练学生适应性沟通。"]: br.add_bullet(doc,t)
    br.add_heading(doc,"2.4 OSCE标准化评价",2)
    for t in ["支持问诊站、查体站、临床思维站、沟通站、急救站和健康宣教站。","量表由客观行为点、关键错误项和整体表现组成；正式成绩以教师锁定版本为准。","保存提问、点击、器材选择、操作顺序、用时、病情变化和沟通表达，形成证据链。","输出总分、分维度表现、扣分证据、标准示范和个性化补练任务。"]: br.add_bullet(doc,t)
    br.add_heading(doc,"2.5 教师内容与教学管理",2)
    for t in ["创建或导入病例，配置适用年级、难度、学习目标和危险信号。","维护知识来源、体征素材、器材映射、沟通脚本和OSCE量表版本。","发布训练任务，设置时长、提示等级、考核模式和可见结果。","查看班级薄弱项、病例难度、评分一致性和内容使用情况。"]: br.add_bullet(doc,t)

    br.add_heading(doc,"三、角色与操作路径",1)
    add_picture(doc,roles,"三类角色协同路径图","课程教师、学生和医学审核教师在内容准备、任务发布、学生训练、评价复盘和持续改进五阶段中的操作。","图2  三类角色操作路径")
    br.add_table(doc,["角色","主要职责","系统中的关键操作"],[
        ["课程教师","定义教学目标并组织教学","选择病例、配置难度与提示、发布任务、查看班级结果"],
        ["学生","完成临床训练并改进能力","问诊、可视化查体、器材检查、诊疗决策、沟通、查看反馈与补练"],
        ["医学审核教师","保证医学正确性和评分公平","审核病例与指南、确认器材规则、锁定量表、复核异常评分"],
        ["系统管理员","维护平台运行和权限","账号、角色权限、日志、版本备份和数据安全配置"],
    ],[1500,3150,4710],8.9)

    br.add_heading(doc,"四、典型教学应用",1)
    scenarios=[
        ("课前自主训练","学生在见习前完成基础病例，教师查看共性遗漏项并调整课堂重点。"),
        ("课堂情境教学","教师投放同一患儿病例，学生分组讨论并比较不同操作导致的病情变化。"),
        ("OSCE考前训练","按照正式考站时长和量表完成练习，隐藏提示，输出逐项证据。"),
        ("低频危重病例","模拟低氧、休克、惊厥等快速变化场景，训练识别、处置顺序与患者安全。"),
        ("传染病与沟通","同时考查接触史、隔离报告和对家长的风险沟通。"),
    ]
    for title,text in scenarios:
        br.add_heading(doc,title,2); br.add_para(doc,text)
    br.add_heading(doc,"建议的武汉大学特色",2)
    for t in ["使用经授权的校本病例、校本OSCE量表和教师标准示范形成核心教学资产。","把“德医双修、知行合一”落实为儿童权益、家长沟通、共同决策等可观察行为。","由儿科学课程教师与附属教学医院教师共同审核病例和量表；未经授权不使用医院名称或诊疗特色。","围绕武汉地区儿童季节性传染病和托幼机构聚集性场景建设专题病例，所有内容标明来源和时间。"]: br.add_bullet(doc,t)

    doc.add_page_break()
    br.add_heading(doc,"五、知识库建设所需资源清单",1)
    br.add_para(doc,"以下材料用于把通用原型转化为可教学、可考核、可追溯的校本系统。T为老师确认启动日，D为比赛最终提交日；如时间冲突，以更早日期为准。真实病例和多媒体素材必须先完成脱敏、授权与医学审核。")
    groups = [
        ("A. 教学目标与课程标准",[
            ["A01 必须","儿科学课程大纲、培养目标、见习/实习要求","建立能力框架和适用年级","T+2；不晚于D-18"],
            ["A02 必须","课程知识点目录、章节结构和重点难点","建立知识标签与检索目录","T+3；不晚于D-17"],
            ["A03 建议","学生常见错误、课程评价和教师教学痛点","设计重点提示与补练任务","T+7；不晚于D-10"],
        ]),
        ("B. 病例与问诊知识",[
            ["B01 必须","首批10—20例脱敏教学病例，至少1例完整示范病例","构建校本病例库和参赛主病例","T+5；不晚于D-15"],
            ["B02 必须","各年龄段标准问诊框架与推荐问法","配置意图识别和问诊评分点","T+5；不晚于D-14"],
            ["B03 必须","病例中患儿与家长分别掌握的信息、情绪和配合度","建立双角色信息差与表达规则","T+6；不晚于D-13"],
            ["B04 建议","教师标准问诊示范或脚本","校准角色回答和示范反馈","T+7；不晚于D-9"],
        ]),
        ("C. 查体、器材与多媒体",[
            ["C01 必须","身体部位—观察动作—可见表面体征映射表","设计患儿模型热点与动作反馈","T+6；不晚于D-13"],
            ["C02 必须","器材—身体部位—操作顺序—可获取信息映射表","实现器材逻辑校验和信息解锁","T+6；不晚于D-13"],
            ["C03 必须","各检查项目标准操作、禁忌和常见错误","建立错误引导和OSCE证据规则","T+6；不晚于D-12"],
            ["C04 条件必须","患儿模型、动作、器材图标、皮疹图片、呼吸音、心音、咳嗽声及授权","建设可视化与多模态查体场景","首批T+7；不晚于D-10"],
            ["C05 必须","体征正常值、年龄分层阈值和危险信号","实现结果解释和安全告警","T+6；不晚于D-12"],
        ]),
        ("D. 传染病与临床规则",[
            ["D01 必须","指定教材、指南、共识清单及版本","建立可追溯医学知识来源","T+5；不晚于D-14"],
            ["D02 必须","儿童传染病年龄、季节、传播和非典型表现要点","建设儿童传染病特点库","T+6；不晚于D-13"],
            ["D03 必须","隔离、报告、院感防控和托幼机构处置流程","建立安全与公共卫生规则","T+6；不晚于D-12"],
            ["D04 必须","年龄/体重相关用药与剂量安全规则的指定来源","建设剂量校验边界；不直接用于临床","T+7；不晚于D-11"],
        ]),
        ("E. 沟通与医学人文",[
            ["E01 必须","典型儿科沟通场景、目标、禁忌表达和评分点","建设沟通任务库","T+5；不晚于D-14"],
            ["E02 必须","家长焦虑、拒检、知情同意、疫苗犹豫等推荐表达","配置角色反应与沟通反馈","T+6；不晚于D-12"],
            ["E03 建议","儿童权益、青春期隐私和坏消息告知规范","建设医学人文考站","T+7；不晚于D-11"],
        ]),
        ("F. OSCE与评价",[
            ["F01 必须","现行儿科OSCE考站任务、评分表和关键错误项","建立正式评价量表","T+3；不晚于D-16"],
            ["F02 必须","评分权重、封顶规则、严重错误和补练映射","实现可解释计分","T+5；不晚于D-14"],
            ["F03 必须","标准答案、可接受表达及争议评分处理原则","支持语义判断与人工复核","T+6；不晚于D-12"],
            ["F04 必须","部位×器材×病例状态测试矩阵及医学验收人","完成正式功能测试与签字","初稿D-12；签字D-7"],
        ]),
        ("G. 治理、平台与参赛材料",[
            ["G01 必须","病例脱敏规范、数据使用范围和伦理/信息部门意见","确定数据使用与部署边界","真实数据导入前"],
            ["G02 必须","知识库内容审核教师名单、专长和审核时间","落实医学审核与争议裁决","T+2；首审D-10"],
            ["G03 必须","最终落地平台、账号权限、网络和设备条件","确定技术路线、权限和部署","T+2；不晚于D-18"],
            ["G04 必须","比赛通知、评分细则、申报模板和截止日期","确定参赛边界与倒排计划","T日；不晚于D-25"],
            ["G05 必须","项目名称、成员、指导教师及署名顺序","统一全部教师汇报和参赛材料","T+1；不晚于D-20"],
            ["G06 条件必须","武汉大学、学院及附属医院名称和视觉标识授权范围","合规呈现校本特色","T+3；不晚于D-15"],
        ]),
    ]
    for heading, rows in groups:
        br.add_heading(doc,heading,2)
        br.add_table(doc,["编号","老师需提供的资料","用途","建议最晚时间"],rows,[1000,3920,3000,1440],7.9)
    br.add_callout(doc,"材料提交原则","不提交真实患儿姓名、病历号、联系方式、精确住址、影像号或任何账号密钥。不能确认授权范围时，先提供目录和脱敏样例，由项目组确认后再批量整理。",fill="FFF0F0",accent=br.RED)
    br.add_heading(doc,"六、老师需要优先确认的事项",1)
    for t in ["确定第一批完整示范病例及其审核教师。","确认身体部位—器材—操作顺序—可获取信息映射表。","提供现行儿科OSCE量表与关键错误项。","确认真实病例、多媒体素材和学校标识的使用边界。","提供比赛通知、评分细则、展示时长和截止日期。"]: br.add_bullet(doc,t)
    br.add_para(doc,"完成上述五项后，即可把功能范围、医学标准和参赛重点锁定，进入正式内容建模和系统实现。",bold=True,color=br.NAVY)
    doc.core_properties.title = "珞珈儿科智训：功能介绍（教师版）"
    doc.core_properties.subject = "系统功能、角色路径与知识库材料需求"
    doc.core_properties.author = "武汉大学儿科学智能体项目组"
    doc.save(FEATURE_OUT)


def build_flow_doc():
    whole = business_flow()
    roles = role_journey()
    student = exam_flow()
    doc = Document()
    setup_doc(doc, "整体业务流程（教师版）")
    add_cover(doc,"珞珈儿科智训：整体业务流程","从教师材料准备到学生训练与教学改进的完整闭环","系统业务流程以教师审核为起点、以学生临床操作为核心、以OSCE证据评价为反馈，并通过版本化知识库持续回到课程改进。")

    br.add_heading(doc,"一、流程总览",1)
    br.add_para(doc,"完整业务由七个连续阶段构成：材料提供、知识加工、医学审核、任务发布、学生训练、OSCE评价和教学改进。任何未经审核的病例、知识或量表都不能进入正式训练或考核。")
    add_picture(doc,whole,"整体业务闭环图","材料提供、知识加工、医学审核、任务发布、学生训练、OSCE评价和教学改进七阶段闭环。","图1  从知识建设到教学改进的整体业务闭环")
    br.add_callout(doc,"流程控制原则","教师拥有病例、知识和评分标准的最终决定权；系统负责执行、记录和反馈，不替代教师进行医学事实确认。")

    br.add_heading(doc,"二、角色协同流程",1)
    add_picture(doc,roles,"三类角色协同路径图","课程教师、学生与医学审核教师在五个阶段的职责与交互。","图2  课程教师、学生与审核教师的协同路径")
    br.add_heading(doc,"2.1 课程教师",2)
    for t in ["提出教学目标，选择病例和适用年级。","配置训练模式、考核模式、时长、难度和提示策略。","发布任务并查看班级完成情况、共性错误与薄弱能力。","根据结果调整课堂重点、病例难度和补练任务。"]: br.add_bullet(doc,t)
    br.add_heading(doc,"2.2 学生",2)
    for t in ["进入教师发布的病例，阅读任务和角色信息。","完成患儿/家长问诊、观察、器材查体、诊疗决策和沟通。","查看OSCE证据报告，理解得分和扣分原因。","进入系统推荐的同类或进阶病例完成补练。"]: br.add_bullet(doc,t)
    br.add_heading(doc,"2.3 医学审核教师",2)
    for t in ["核验病例事实、指南来源、器材规则、危险信号和评分量表。","确认内容适用年龄、教学阶段及使用边界。","锁定正式版本，复核异常结果和争议评分。","批准知识更新、病例退役和量表版本切换。"]: br.add_bullet(doc,t)

    br.add_heading(doc,"三、知识库建设流程",1)
    stages = [
        ("3.1 材料接收","老师按材料清单提供课程大纲、病例、指南、体征、沟通场景和OSCE量表。项目组只接收脱敏、可授权、可追溯材料。"),
        ("3.2 脱敏与格式检查","检查身份信息、文件完整性、来源、版本、授权范围和适用年级；不合格材料返回补充。"),
        ("3.3 结构化拆解","把材料拆分为病例事实、角色信息、问诊主题、身体部位、器材规则、体征、危险信号、沟通行为点和评分项。"),
        ("3.4 关系与规则映射","建立“年龄—病例阶段—部位—器材—操作顺序—可获取信息—反馈—评分点”映射，并配置病情状态迁移。"),
        ("3.5 医学审核","至少由指定儿科学教师核对内容；正式考核病例建议双人审核。争议内容不得上线。"),
        ("3.6 版本发布","记录内容版本、审核人、发布日期和复核日期；训练与考核引用固定版本，更新不能影响已开始的考站。"),
        ("3.7 运行监测与更新","收集错误案例、低命中问题和评分争议，教师确认后形成下一版本；过期指南及时下线。"),
    ]
    n = br.create_decimal_numbering(doc)
    for title,text in stages:
        br.add_heading(doc,title,2); br.add_para(doc,text)
    br.add_table(doc,["质量关口","检查内容","不通过时处理"],[
        ["隐私关","是否包含可识别患儿身份信息","拒绝入库，返回脱敏"],
        ["来源关","是否有教材、指南或教师确认来源","标记待审，不用于正式训练"],
        ["医学关","病例事实、操作位置、器材规则是否正确","退回修改并重新审核"],
        ["评价关","行为点、分值、关键错误是否完整一致","不得用于正式考核"],
        ["版本关","审核人、版本和复核日期是否齐全","不得发布或覆盖旧版本"],
    ],[1450,4200,3710],8.8)

    br.add_heading(doc,"四、学生训练业务流程",1)
    add_picture(doc,student,"学生一次训练完整路径图","学生从进入病例到双角色问诊、可视化查体、器材检查、临床决策、沟通和评价补练的七步流程。","图3  学生一次训练的完整路径")
    br.add_heading(doc,"4.1 进入病例与问诊",2)
    br.add_para(doc,"学生先确认患儿年龄、主诉和任务要求，再选择与患儿或家长交流。系统按角色事实范围回答，并记录学生是否覆盖主诉、现病史、出生喂养、生长发育、接种、接触和危险症状。")
    br.add_heading(doc,"4.2 可视化观察与器材查体",2)
    br.add_para(doc,"学生点击患儿身体部位时，系统只展示可直接观察的表面体征或动作。例如点击口腔后患儿张口、吐舌。需要器材才能获得的信息必须经过校验：选择器材→定位身体部位→满足准备和操作顺序→展示结果。错误器材、错误部位或越序操作不能解锁深层信息。")
    br.add_table(doc,["学生动作","系统校验","反馈结果"],[
        ["直接点击身体部位","部位是否可观察、患儿是否配合","播放动作并展示表面体征"],
        ["选择体温计测温","器材、测量部位和病例状态是否匹配","展示体温及异常提示"],
        ["使用听诊器听诊胸部","听诊器、胸部区域和听诊顺序是否正确","播放对应呼吸音并记录部位"],
        ["血压计用于错误部位","器材与部位不匹配","不显示血压，提示重新选择"],
        ["遗漏关键查体","考站结束时检查证据完整性","量表扣分并推荐针对性补练"],
    ],[2600,3350,3410],8.7)
    br.add_heading(doc,"4.3 临床决策与病情演变",2)
    br.add_para(doc,"学生提交诊断、检查或处置后，系统读取当前病例状态和安全规则。正确处置可使生命体征稳定，延误或错误处置可使风险等级升高。所有状态迁移由教师预设，不由模型临时编造。")
    br.add_heading(doc,"4.4 沟通、评价与补练",2)
    br.add_para(doc,"学生向家长解释判断、风险和下一步计划。系统记录共情、通俗表达、共同决策和安全网告知等行为。考站结束后，固定量表读取行为证据生成分项结果；自然语言模型仅辅助识别表达，不独立决定最终成绩。")

    br.add_heading(doc,"五、OSCE评价流程",1)
    for s in ["收集全过程证据：提问、点击、器材、部位、顺序、用时、诊疗决策和沟通表达。","把证据映射到教师锁定的行为点和关键错误项。","先计算客观行为分，再应用封顶规则、严重错误和整体表现分。","生成总分、能力维度、扣分证据和标准示范。","根据薄弱点匹配同类病例或进阶病例，形成补练路径。","教师抽查异常评分和争议记录，必要时人工复核并更新规则。"]: br.add_number(doc,s,n)
    br.add_callout(doc,"评价边界","模型可以帮助理解学生的自然语言，但不能修改病例事实、量表权重或关键错误项；正式成绩必须可回查到具体行为证据。",fill=br.PALE_GOLD,accent=br.GOLD)

    br.add_heading(doc,"六、异常与安全流程",1)
    br.add_table(doc,["异常情形","系统处理","教师/管理员处理"],[
        ["学生询问病例未提供的信息","明确提示“当前病例未提供”，不补写体征","审核是否需要补充病例事实"],
        ["器材或部位选择错误","阻止结果展示，给予分级纠错","确认错误提示是否符合教学要求"],
        ["危险处置或剂量异常","触发患者安全提示并记录关键错误","审核规则、必要时停止病例使用"],
        ["知识来源过期","标记到期并停止进入正式考核版本","更新指南并重新审核发布"],
        ["评分争议","保留证据、量表版本和计算过程","审核教师复核并记录裁决"],
        ["系统或模型不可用","保存当前进度；使用固定规则或暂停考核","恢复服务后核验记录，不能静默丢分"],
    ],[2200,3500,3660],8.5)

    br.add_heading(doc,"七、老师确认与上线条件",1)
    for t in ["至少完成1个全流程病例，病例事实、身体部位、器材规则和病情演变均经教师审核。","现行OSCE量表、关键错误项和评分权重已锁定版本。","所有图片、音频、动画和器材素材来源合法，使用范围明确。","部位×器材×病例状态测试矩阵全部通过，错误组合不得泄露深层信息。","至少完成5名目标学生试用和1轮教师评分一致性校准。","数据脱敏、权限、审计、备份和故障处理流程完成确认。"]: br.add_bullet(doc,t)
    br.add_heading(doc,"八、与知识库材料清单的对应关系",1)
    br.add_para(doc,"知识库建设所需全部资料已列入《功能介绍（教师版）》第五部分，按教学目标、病例问诊、查体器材、传染病规则、沟通人文、OSCE评价及治理平台七类组织。老师可优先提供：完整示范病例、器材—部位映射表、现行OSCE量表、审核教师名单和比赛评分细则。")
    br.add_callout(doc,"建议启动顺序","先锁定1个参赛主病例和1套量表，再完成身体部位与器材规则，随后补充多媒体素材和其他病例。这样能够尽快形成医学正确、流程完整、可复用的第一版。")
    remove_trailing_empty_paragraph(doc)
    doc.core_properties.title = "珞珈儿科智训：整体业务流程（教师版）"
    doc.core_properties.subject = "角色操作路径、知识库建设、学生训练与OSCE评价流程"
    doc.core_properties.author = "武汉大学儿科学智能体项目组"
    doc.save(FLOW_OUT)


if __name__ == "__main__":
    build_feature_doc()
    build_flow_doc()
    print(FEATURE_OUT)
    print(FLOW_OUT)
