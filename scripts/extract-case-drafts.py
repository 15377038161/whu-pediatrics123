"""把病例原文抽取到本地忽略目录，供教师核对；绝不自动发布为学生病例事实。"""

from hashlib import sha256
from pathlib import Path
import json
import re
import os
import subprocess
import shutil
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from pypdf import PdfReader

ROOT = Path('knowledge/儿科常见病')
CATALOG = Path('src/domain/teacher-case-catalog.ts')
OUTPUT = Path('output/case-drafts')
DOCUMENT_EXTENSIONS = {'.doc', '.docx', '.pdf', '.pptx'}
SECTION_PATTERN = re.compile(r'^(主诉|现病史|既往史|个人史|家族史|体格检查|查体|辅助检查|初步诊断|诊断|鉴别诊断|治疗|诊疗经过|出院诊断)[：:\s]')


def read_catalog():
    source = CATALOG.read_text(encoding='utf-8')
    start = source.index('= [') + 2
    return json.loads(source[start:source.rindex(']') + 1])


def read_docx(file):
    document = Document(file)
    lines = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            line = ' | '.join(cell.text.strip().replace('\n', ' / ') for cell in row.cells)
            if line.strip(' |'):
                lines.append(line)
    return '\n'.join(lines), len(lines)


def read_pdf(file):
    reader = PdfReader(str(file))
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or '').strip()
        if text:
            pages.append(f'[第{number}页]\n{text}')
    return '\n\n'.join(pages), len(reader.pages)


def read_pptx(file):
    pages = []
    with zipfile.ZipFile(file) as archive:
        slides = sorted((name for name in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', name)), key=lambda name: int(re.search(r'slide(\d+)', name).group(1)))
        for index, slide in enumerate(slides, start=1):
            root = ET.fromstring(archive.read(slide))
            lines = [''.join(node.itertext()) for node in root.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t')]
            pages.append(f'[幻灯片{index}]\n' + '\n'.join(lines))
    return '\n\n'.join(pages), len(pages)


def extract_file(file):
    relative = file.as_posix()
    record = {'source': relative, 'sha256': sha256(file.read_bytes()).hexdigest(), 'status': 'unsupported', 'unitCount': 0, 'text': '', 'candidateHeadings': []}
    try:
        if file.suffix.lower() == '.docx':
            record['text'], record['unitCount'] = read_docx(file)
        elif file.suffix.lower() == '.pdf':
            record['text'], record['unitCount'] = read_pdf(file)
        elif file.suffix.lower() == '.pptx':
            record['text'], record['unitCount'] = read_pptx(file)
        else:
            converter = os.environ.get('LIBREOFFICE_BIN') or shutil.which('soffice')
            if not converter and Path('C:/Program Files/LibreOffice/program/soffice.com').exists():
                converter = 'C:/Program Files/LibreOffice/program/soffice.com'
            target = OUTPUT / 'converted' / record['sha256']
            converted = target / (file.stem + '.docx')
            if not converted.exists() and converter:
                target.mkdir(parents=True, exist_ok=True)
                profile = (OUTPUT / 'libreoffice-profile').resolve().as_uri()
                subprocess.run([converter, f'-env:UserInstallation={profile}', '--headless', '--convert-to', 'docx', '--outdir', str(target.resolve()), str(file.resolve())], capture_output=True, timeout=90, check=True)
            if not converted.exists():
                record['status'] = 'legacy-doc-needs-conversion'
                return record
            record['text'], record['unitCount'] = read_docx(converted)
        record['status'] = 'text-extracted' if record['text'].strip() else 'no-text-ocr-needed'
        record['candidateHeadings'] = [line[:80] for line in record['text'].splitlines() if SECTION_PATTERN.match(line.strip())][:40]
    except Exception as error:
        record['status'] = 'extract-error'
        record['errorType'] = type(error).__name__
    return record


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    cases = read_catalog()
    counts = {}
    for item in cases:
        folder = Path(item['folder']) if item['folder'] else None
        documents = sorted((p for p in folder.rglob('*') if p.is_file() and p.suffix.lower() in DOCUMENT_EXTENSIONS), key=lambda p: str(p)) if folder else []
        records = [extract_file(file) for file in documents]
        for record in records:
            counts[record['status']] = counts.get(record['status'], 0) + 1
        draft = {'caseId': item['id'], 'name': item['name'], 'source': item['source'], 'catalogStatus': item['status'],
                 'reviewStatus': 'unreviewed', 'mediaCount': item['mediaCount'], 'documents': records}
        (OUTPUT / f"{item['id']}.json").write_text(json.dumps(draft, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'cases': len(cases), 'documents': sum(counts.values()), 'statuses': counts, 'output': str(OUTPUT)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
