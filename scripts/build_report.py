from pathlib import Path
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)
OUT = DOCS / "珞珈儿科智训_智能体建设汇报.docx"

NAVY = "0A3556"
BLUE = "176FA9"
CYAN = "42B9C8"
GOLD = "C39532"
INK = "1B2E3B"
MUTED = "60727E"
LINE = "D9E4E9"
LIGHT = "F3F7F9"
PALE_BLUE = "EAF4F8"
PALE_GOLD = "FFF7E7"
WHITE = "FFFFFF"
RED = "A94442"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_border(cell, color=LINE, size=6):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:color"), color)


def set_run_font(run, name="Microsoft YaHei", size=None, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)


def add_para(doc, text="", size=10.5, color=INK, bold=False, italic=False,
             align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=8, line=1.333, style=None):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.left_indent = Inches(0.375 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.194)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    r = p.add_run(text)
    set_run_font(r, size=10.3)
    return p


def create_decimal_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType"); multi.set(qn("w:val"), "singleLevel"); abstract.append(multi)
    level = OxmlElement("w:lvl"); level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1"); level.append(start)
    fmt = OxmlElement("w:numFmt"); fmt.set(qn("w:val"), "decimal"); level.append(fmt)
    text = OxmlElement("w:lvlText"); text.set(qn("w:val"), "%1."); level.append(text)
    jc = OxmlElement("w:lvlJc"); jc.set(qn("w:val"), "left"); level.append(jc)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab"); tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), "540"); tabs.append(tab); ppr.append(tabs)
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "540"); ind.set(qn("w:hanging"), "280"); ppr.append(ind)
    level.append(ppr); abstract.append(level); numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId"); abstract_ref.set(qn("w:val"), str(abstract_id)); num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_number(doc, text, num_id):
    p = doc.add_paragraph()
    ppr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0"); num_pr.append(ilvl)
    num = OxmlElement("w:numId"); num.set(qn("w:val"), str(num_id)); num_pr.append(num)
    ppr.append(num_pr)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    r = p.add_run(text)
    set_run_font(r, size=10.3)
    return p


def add_callout(doc, label, text, fill=PALE_BLUE, accent=BLUE):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), "22"); left.set(qn("w:color"), accent)
    borders.append(left); p_pr.append(borders)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "140"); spacing.set(qn("w:after"), "140")
    p_pr.append(spacing)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(label + "  ")
    set_run_font(r, size=10.5, color=accent, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    add_para(doc, "", after=3)


def set_last_image_alt(doc, title, description):
    inline = doc.inline_shapes[-1]._inline
    inline.docPr.set("title", title)
    inline.docPr.set("descr", description)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size={1:16, 2:13, 3:11.5}[level], color={1:NAVY,2:BLUE,3:"245B78"}[level], bold=True)
    return p


def add_table(doc, headers, rows, widths, font_size=8.9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=9, color=WHITE, bold=True)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            if row_idx % 2:
                set_cell_shading(cells[idx], "F8FAFB")
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.18
            if idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK, bold=(idx == 0 and len(str(value)) < 8))
    set_table_geometry(table, widths)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            set_cell_border(cell)
    add_para(doc, "", after=3)
    return table


def font(size, bold=False):
    path = Path("C:/Windows/Fonts/msyh.ttc")
    if not path.exists():
        path = Path("C:/Windows/Fonts/simhei.ttf")
    return ImageFont.truetype(str(path), size=size, index=0) if path.exists() else ImageFont.load_default()


def rounded(draw, box, fill, outline=None, width=2, radius=18):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def build_architecture_image():
    img = Image.new("RGB", (1800, 980), "#F4F8FA")
    d = ImageDraw.Draw(img)
    title = font(44, True); h = font(29, True); body = font(23); small = font(20)
    d.text((70, 42), "珞珈儿科智训整体架构", font=title, fill="#0A3556")
    d.text((70, 100), "教学资源由教师审核，智能引擎负责模拟、演变与评价，输出服务训练和竞赛展示。", font=body, fill="#60727E")
    # input layer
    rounded(d, (70, 175, 1730, 315), "#FFFFFF", "#D9E4E9", 3)
    d.text((100, 198), "输入层", font=h, fill="#176FA9")
    inputs = ["学生自然语言", "查体与处置选择", "患儿/家长角色", "教师配置"]
    for i, text in enumerate(inputs):
        x = 300 + i * 345
        rounded(d, (x, 205, x+285, 280), "#EAF4F8")
        d.text((x+28, 226), text, font=body, fill="#164E70")
    # engine layer
    rounded(d, (70, 365, 1730, 615), "#0A3556")
    d.text((100, 392), "核心智能层", font=h, fill="#78D6DF")
    engines = [("双角色模拟引擎","年龄化表达\n信息差与情绪状态"),("动态病例引擎","病情阶段\n生命体征与反馈"),("可解释评价引擎","量表计分\n证据与补练建议")]
    for i,(name,desc) in enumerate(engines):
        x=290+i*480
        rounded(d,(x,420,x+390,565),"#FFFFFF")
        d.text((x+28,445),name,font=h,fill="#0A3556")
        lines=desc.split("\n")
        d.text((x+28,500),lines[0]+" · "+lines[1],font=small,fill="#60727E")
        if i<2:
            d.line((x+405,492,x+462,492),fill="#66C7D0",width=5)
            d.polygon([(x+462,492),(x+445,480),(x+445,504)],fill="#66C7D0")
    # knowledge base layer
    d.text((80, 665), "校本知识与规则层", font=h, fill="#176FA9")
    bases=["问诊与体征模拟库","儿童传染病特点库","沟通场景专项库","OSCE标准化考核库"]
    for i,name in enumerate(bases):
        x=70+i*420
        rounded(d,(x,710,x+385,810),"#FFFFFF","#BFD9E4",2)
        d.text((x+24,744),name,font=body,fill="#0A3556")
    # output strip
    rounded(d,(70,860,1730,940),"#FFF5DF","#E8D098",2)
    outputs="输出：角色回答  ·  体征反馈  ·  病情变化  ·  OSCE评分  ·  过程证据  ·  个性化补练"
    d.text((118,884),outputs,font=body,fill="#6C5117")
    path=ASSETS/"architecture.png"
    img.save(path,quality=95)
    return path


def build_workflow_image():
    img=Image.new("RGB",(1800,520),"#FFFFFF")
    d=ImageDraw.Draw(img); h=font(28,True); body=font(20)
    steps=[("01","学生输入","问诊/查体/决策"),("02","意图识别","识别主题与操作"),("03","病例状态","读取角色和阶段"),("04","专业校验","四库与安全规则"),("05","生成反馈","回答/体征/演变"),("06","证据评价","评分与补练")]
    colors=["#EAF4F8","#EAF4F8","#E9F7F3","#E9F7F3","#FFF6E4","#FFF6E4"]
    for i,(num,name,desc) in enumerate(steps):
        x=40+i*290
        rounded(d,(x,130,x+245,360),colors[i],"#C9DDE6",2,20)
        d.text((x+25,155),num,font=h,fill="#42AFC0")
        d.text((x+25,220),name,font=h,fill="#0A3556")
        d.text((x+25,278),desc,font=body,fill="#60727E")
        if i<5:
            d.line((x+250,245,x+285,245),fill="#8CB4C5",width=5)
            d.polygon([(x+285,245),(x+270,234),(x+270,256)],fill="#8CB4C5")
    path=ASSETS/"workflow.png";img.save(path,quality=95);return path


def setup_styles(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    for level, size, color, before, after in ((1,16,NAVY,18,10),(2,13,BLUE,12,6),(3,11.5,"245B78",8,4)):
        s=doc.styles[f"Heading {level}"]
        s.font.name="Microsoft YaHei";s._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
        s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor.from_string(color)
        s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after);s.paragraph_format.keep_with_next=True


def setup_header_footer(doc):
    sec=doc.sections[0]
    hp=sec.header.paragraphs[0]
    hp.alignment=WD_ALIGN_PARAGRAPH.LEFT
    r=hp.add_run("珞珈儿科智训  |  智能体大赛建设汇报")
    set_run_font(r,size=8.5,color=MUTED)
    fp=sec.footer.paragraphs[0]
    fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=fp.add_run("武汉大学儿科学教学智能体  ·  ")
    set_run_font(r,size=8,color=MUTED)
    add_field(fp,"PAGE")


def cover(doc):
    add_para(doc,"武汉大学儿科学 · 智能体大赛项目",size=11,color=GOLD,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=28)
    add_para(doc,"珞珈儿科智训",size=31,color=NAVY,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=10)
    add_para(doc,"儿童全周期临床胜任力训练与OSCE考核智能体",size=16,color=BLUE,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=28)
    add_para(doc,"让儿科临床训练可重复、可解释、可改进",size=12,color=MUTED,italic=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=46)
    add_callout(doc,"汇报结论", "本项目以儿科学专业能力为中心，以患儿—家长双角色、动态病情演变和可解释OSCE评价为三项核心创新，贯通问诊、查体、诊断、沟通与考核，形成具有武汉大学校本特色的竞赛型教学智能体。",fill=PALE_BLUE)
    add_para(doc,"交付版本：竞赛演示原型 V0.1",size=10,color=MUTED,align=WD_ALIGN_PARAGRAPH.CENTER,before=55,after=5)
    add_para(doc,f"编制日期：{date.today().strftime('%Y年%m月%d日')}",size=10,color=MUTED,align=WD_ALIGN_PARAGRAPH.CENTER,after=5)
    add_para(doc,"文档状态：供教师论证、材料准备与参赛决策使用",size=10,color=MUTED,align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def build_document():
    arch=build_architecture_image();flow=build_workflow_image()
    doc=Document();setup_styles(doc);setup_header_footer(doc);cover(doc)

    add_heading(doc,"汇报摘要",1)
    add_para(doc,"本项目拟建设“珞珈儿科智训”智能体，服务儿科学专业学生的临床胜任力训练与OSCE考核。项目不定位为一般医学问答工具，而是通过一个可演变的虚拟病例，将学生的语言输入、查体选择、临床决策和沟通行为转化为可追踪、可评分、可复训的教学过程。")
    add_callout(doc,"核心价值", "解决真实患儿难以重复配合、低频危重病例接触不足、儿科三方沟通训练缺位、传统OSCE组织成本高且反馈滞后的问题。")
    add_heading(doc,"本次交付范围",2)
    add_bullet(doc,"一份面向教师的正式汇报文档：完整拆解总体架构、实现逻辑、建设路径、风险边界和材料需求。")
    add_bullet(doc,"一套可交互演示Demo：支持教师手动问诊、查体、决策、沟通，也支持一键完整演示。")
    add_bullet(doc,"一个可直接展示的交付包：无需API密钥或服务器即可演示，附使用说明与验收记录。")
    add_heading(doc,"当前已经具备的演示能力",2)
    add_table(doc,["环节","教师可见内容","系统执行逻辑"],[
        ["输入","学生向患儿或家长提问；选择查体、诊断和沟通方案","识别问诊主题和操作类型，写入本次考站状态"],
        ["处理","患儿/家长回答、体征解锁、生命体征变化","按角色、病例阶段和专业规则匹配反馈"],
        ["输出","OSCE总分、六维雷达图、逐项证据和补练建议","量表计分，绑定已发生的行为证据并识别薄弱项"],
    ],[1200,3650,4510],9.2)

    add_heading(doc,"一、建设背景与参赛定位",1)
    add_heading(doc,"1.1 儿科学临床教学的典型痛点",2)
    for t in [
        "真实患儿年龄小、表达有限、情绪波动大，难以像成人标准化病人一样稳定重复配合。",
        "儿科病史不仅来自患儿，还依赖家长补充；学生容易忽略出生史、喂养史、生长发育史和预防接种史。",
        "病情变化快，急危重症和传染病处置具有低频、高风险特点，学生真实练习机会不足。",
        "传统OSCE需要大量教师、标准化病人、场地和排期，考后往往只能得到分数，缺少全过程证据。",
        "通用大模型善于回答知识问题，但不能天然保证病例一致性、评分公平性和医学安全。",
    ]: add_bullet(doc,t)
    add_heading(doc,"1.2 竞赛作品定位",2)
    add_para(doc,"参赛作品以“儿童全周期临床胜任力训练”为单一主线，用一个高完成度病例贯通四个专业资源库和三个核心引擎。竞赛现场不以功能菜单数量取胜，而以完整的输入—逻辑—输出闭环、儿科专业差异和教师可控的评价证据形成辨识度。")
    add_callout(doc,"一句话介绍", "AI同时扮演患儿与家长，学生的问诊、查体和处置会推动病情变化，最终依据教师确认的OSCE量表生成证据化反馈。",fill=PALE_GOLD,accent=GOLD)
    add_heading(doc,"1.3 建设目标",2)
    for t in ["专业目标：覆盖儿科问诊、查体、临床思维、传染病特点、沟通与患者安全。","教学目标：支持训练模式、考核模式、教师复盘和个性化补练。","竞赛目标：形成可现场稳定演示、可量化说明价值、可继续扩展的完整作品。","治理目标：确保病例脱敏、知识可追溯、评分可解释、模型有边界。"]: add_bullet(doc,t)

    add_heading(doc,"二、整体架构拆解",1)
    add_para(doc,"系统由输入层、核心智能层、校本知识与规则层、输出层组成。教师负责病例、量表、指南和教学目标；系统负责多角色模拟、病例状态推进、规则校验和过程评价。")
    doc.add_picture(str(arch),width=Inches(6.45))
    set_last_image_alt(doc,"珞珈儿科智训整体架构","输入层、双角色模拟引擎、动态病例引擎、可解释评价引擎、四个校本资源库和输出层的关系图。")
    p=doc.paragraphs[-1];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(6)
    add_para(doc,"图1  珞珈儿科智训整体架构（当前Demo采用本地规则模拟核心链路）",size=8.5,color=MUTED,align=WD_ALIGN_PARAGRAPH.CENTER)
    add_heading(doc,"2.1 输入层",2)
    add_bullet(doc,"自然语言输入：学生向患儿或家长提出问诊问题。")
    add_bullet(doc,"结构化操作：学生选择查体项目、检查、初步诊断、处置和沟通表达。")
    add_bullet(doc,"教师配置：病例难度、时间、评分项、提示策略、知识版本和安全边界。")
    add_heading(doc,"2.2 核心智能层",2)
    add_bullet(doc,"双角色模拟引擎：管理患儿和家长的信息差、语言能力、情绪与配合程度。")
    add_bullet(doc,"动态病例引擎：依据时间、已获信息和学生操作更新病情阶段与生命体征。")
    add_bullet(doc,"可解释评价引擎：使用教师量表计分，将每个分值绑定到实际行为证据。")
    add_heading(doc,"2.3 校本知识与规则层",2)
    add_bullet(doc,"病例知识：病史、体征、检查、诊断、治疗、病情演变和教学目标。")
    add_bullet(doc,"专业规则：年龄相关正常值、药物剂量、危险信号、隔离与报告要求。")
    add_bullet(doc,"评价规则：OSCE行为点、整体表现、严重错误项和补练映射。")
    add_heading(doc,"2.4 输出层",2)
    add_bullet(doc,"即时输出：角色回答、体征反馈、病情变化、风险提示。")
    add_bullet(doc,"考后输出：总分、能力维度、扣分证据、标准示范和个性化补练。")
    add_bullet(doc,"教师输出：班级薄弱项、病例难度、评分一致性和内容版本记录。")

    add_heading(doc,"三、四个专业资源库",1)
    modules=[
        ("3.1 儿科问诊与体征模拟库",["按新生儿、婴儿、幼儿、学龄儿童和青春期设置表达能力与查体配合度。","覆盖主诉、现病史、出生史、喂养史、生长发育史、预防接种史、接触史与家族史。","建设可视化患儿交互模型：学生可自主点击口腔、胸部、腹部、皮肤等身体部位，触发符合年龄和病例状态的动作或表面体征反馈，例如点击嘴巴后患儿张口、吐舌，显示可直接观察的信息。","体检器材必须由学生主动选择并拖放或点击到正确部位；听诊器、血压计、体温计等工具与部位匹配成功后才解锁呼吸音、血压、体温等深层信息。","病例结构包含部位热点、器材—部位规则、表面体征、深层病情信息、错误引导、标准答案、阶段变化和教师点评。"]),
        ("3.2 儿童传染病特点库",["按年龄、季节、传播途径、临床综合征和隔离要求组织。","突出儿童非典型表现、年龄依赖性、疫苗接种、剂量安全和快速恶化风险。","训练流行病学史、传染病报告、院感防控、家庭与托幼机构健康教育。","所有医学内容记录指南名称、版本、发布日期、审核教师和失效复核日期。"]),
        ("3.3 儿科沟通场景专项库",["覆盖哭闹拒检、焦虑家长、抗菌药诉求、疫苗犹豫、知情同意、坏消息告知和青春期隐私。","评价共情、结构、通俗程度、风险解释、共同决策和儿童权益保护。","同一病例可设置不同家长性格与认知水平，考查学生适应性沟通。","沟通评分以可观察行为点为主，模型仅提供语义辅助，不直接决定全部分值。"]),
        ("3.4 儿科OSCE标准化考核库",["支持问诊站、查体站、临床思维站、沟通站、急救站和健康宣教站。","采用客观行为点、关键错误项和整体表现三部分评分。","保存学生提问、选择、顺序、用时和病情变化，形成完整证据链。","训练模式允许分级提示；正式考核模式隐藏提示并锁定量表版本。"]),
    ]
    for title,items in modules:
        add_heading(doc,title,2)
        for item in items:add_bullet(doc,item)

    add_heading(doc,"四、三个核心引擎与实现逻辑",1)
    add_heading(doc,"4.1 患儿—家长双角色模拟引擎",2)
    add_para(doc,"每个角色拥有独立的事实范围、表达能力、情绪状态和信息可信度。患儿提供主观感受和有限表达，家长补充时间线、既往史和照护观察。系统根据学生当前交流对象输出不同答案，避免两个角色说出完全相同的信息。")
    add_heading(doc,"4.2 动态病情演变引擎",2)
    add_para(doc,"病例被表示为有限状态：初始评估、危险信号出现、处置后稳定或延误后恶化。学生完成关键操作后触发状态迁移，例如识别低氧并吸氧后血氧上升；忽视危险信号则风险等级提高。正式版的状态迁移必须由教师预先确认，不由模型随意编造。")
    add_heading(doc,"4.3 可解释OSCE评价引擎",2)
    add_para(doc,"评分先读取结构化行为证据，再套用教师量表计算。自然语言模型只负责识别学生表达是否覆盖某个行为点，最终分值由固定权重、封顶规则和关键错误项共同决定。每项扣分均能回指到“学生做了什么或没有做什么”。")
    add_heading(doc,"4.4 可视化交互式患儿查体子系统（正式建设计划）",2)
    add_callout(doc,"关键升级", "模拟问诊系统不能停留在对话问答。正式版将把患儿模型作为主要操作界面，让学生按照真实临床顺序完成观察、选择器材、定位部位、执行检查和解释结果。当前Demo仅验证“主动操作后开放体征”的规则闭环，不代表该可视化子系统已经完成。",fill=PALE_GOLD,accent=GOLD)
    add_table(doc,["交互层级","学生操作","系统逻辑与反馈"],[
        ["基础观察","点击口腔、胸部、腹部、皮肤等身体部位","播放张口、吐舌、转身、抬臂等动作；仅显示肉眼可见的表面体征"],
        ["器材调用","从侧边工具栏选择听诊器、血压计、体温计等","显示器材已选中、可用部位和操作步骤，不直接泄露检查结果"],
        ["专业校验","将器材用于具体身体部位","同时匹配“病例状态＋器材＋部位＋操作顺序”后，才解锁呼吸音、血压、体温等深层信息"],
        ["错误引导","器材或部位选择错误、遗漏准备步骤","给予分级提示：状态提示→方向性提示→标准操作示范；记录错误次数和用时"],
        ["教学评价","完成或结束一次检查","保存点击轨迹、器材选择、部位匹配、操作顺序和结果解释，映射到OSCE评分证据"],
    ],[1500,2900,4960],8.6)
    add_para(doc,"核心数据规则采用“病例阶段—年龄—身体部位—观察方式—器材—可见信息—反馈动画—评分点”的结构化映射。没有使用对应器材时，系统不得展示深层病理数据；器材与部位不匹配时，不触发检查结果，只提供纠错提示。所有动作、检查位置和结果解释由儿科学教师审核后进入正式病例。")
    add_heading(doc,"4.5 一次交互的完整处理链",2)
    doc.add_picture(str(flow),width=Inches(6.45));set_last_image_alt(doc,"六步处理链","学生输入、意图识别、病例状态、专业校验、生成反馈和证据评价六步流程图。");doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc,"图2  从学生输入到评价输出的六步处理链",size=8.5,color=MUTED,align=WD_ALIGN_PARAGRAPH.CENTER)
    steps=["接收学生的问诊语句或结构化操作。","识别问诊主题、临床操作类型和当前交流角色。","读取病例事实、当前病情阶段和已披露信息。","调用四库内容及危险信号、年龄、剂量等专业规则。","生成符合角色和阶段的回答、体征或生命体征变化。","记录过程证据，按量表计分并生成补练建议。"]
    workflow_num = create_decimal_numbering(doc)
    for s in steps:add_number(doc,s,workflow_num)

    add_heading(doc,"五、演示Demo设计与已实现功能",1)
    add_heading(doc,"5.1 演示病例",2)
    add_para(doc,"演示病例为“3岁患儿发热、咳嗽伴气促”。学生需要完成重点问诊、呼吸系统查体、低氧风险识别、初步处置和家长沟通。该病例能够在约5分钟内展示四库三引擎的完整闭环。病例和评分为虚构演示数据，不构成诊疗建议。")
    add_heading(doc,"5.2 手动操作路径",2)
    demo_num = create_decimal_numbering(doc)
    for s in ["选择患儿母亲或患儿小宇，通过输入框完成问诊。","点击查体项目，观察系统按需开放体征。","提交初步诊断与首要处置，查看生命体征变化。","选择向焦虑家长的沟通表达并获得反馈。","完成考站，查看六维雷达图、逐项证据和个性化补练。"]:add_number(doc,s,demo_num)
    add_heading(doc,"5.3 一键演示路径",2)
    add_para(doc,"教师点击右上角“一键演示完整流程”，系统自动完成标准问诊、重点查体、正确处置和沟通，随后跳转OSCE报告。该模式用于竞赛现场快速、稳定地呈现完整价值；任何阶段都可以由教师手动接管。")
    add_heading(doc,"5.4 当前技术实现",2)
    add_table(doc,["组件","当前Demo","正式版本建议"],[
        ["界面","HTML、CSS、JavaScript，纯前端运行","保持网页端，接入学校统一身份认证"],
        ["意图识别","本地关键词与规则","知识库约束的大模型＋规则校验"],
        ["病例状态","浏览器内存中的有限状态","后端病例状态机＋数据库审计"],
        ["OSCE评分","固定权重与行为证据","教师版本化量表＋语义判定＋人工抽查"],
        ["数据","虚构演示病例","脱敏校本病例、授权体征素材、指南库"],
        ["部署","双击打开，无需联网","可部署至扣子、独立Web或校内平台"],
    ],[1600,3300,4460],8.8)
    add_heading(doc,"5.5 当前Demo与计划功能边界",2)
    add_para(doc,"当前交付Demo已经实现问诊、查体项目选择、器材/动作规则的简化逻辑、病情演变与OSCE报告，适合教师快速理解完整闭环。可点击患儿身体部位、动作动画、器材拖放和精细部位匹配属于正式建设计划，需在取得教师确认的检查规范、部位映射及合法素材后开发，不应在现阶段汇报中表述为已完成。")

    add_heading(doc,"六、武汉大学专有特色设计",1)
    add_callout(doc,"设计原则", "学校特色必须进入病例、评价标准和培养过程，而不是只使用校徽、校色和校园图案。")
    features=[
        ("珞珈儿科校本病例库","由武汉大学儿科学教师及附属教学医院提供脱敏病例，记录来源科室、适用年级、教学目标和审核人，形成不可简单复制的教学资产。"),
        ("珞珈儿科能力评价标准","将课程目标、见习实习要求和OSCE量表统一，保证日常训练与考核使用同一能力标准。"),
        ("“德医双修、知行合一”人文考站","围绕儿童权益、家长焦虑、青春期隐私、重大疾病告知和共同决策，把医学人文变成可观察、可评分的行为。"),
        ("附属医院协同审核","人民医院、中南医院等教学力量可共同建设病例和审核量表；未获得正式授权前，不在系统中虚构医院诊疗特色。"),
        ("武汉地区儿童健康专题","可依据教师确认的疾控资料建设季节性传染病、托幼机构聚集性疫情和特殊天气应急场景，并保持时间与来源标注。"),
        ("临床问题到科研问题","病例结束后将诊疗疑问转化为PICO问题，训练循证检索、证据评价和跨学科创新思维。"),
    ]
    for title,text in features:
        add_heading(doc,title,2);add_para(doc,text)
    add_para(doc,"武汉大学医学教育特色参考：学校医学教育强调“德医双修、知行合一”、临床能力、创新思维、模块化整合课程与跨学科培养。本项目将这些理念落实为病例、任务和评价行为。",size=9.2,color=MUTED,italic=True)

    add_heading(doc,"七、竞赛创新点与差异化",1)
    add_table(doc,["创新点","传统方案","本项目差异"],[
        ["双角色标准化病人","单一AI患者或固定脚本","患儿与家长拥有不同信息、情绪和表达能力"],
        ["动态病例","问答结果不影响病例","学生操作驱动病情和生命体征变化"],
        ["可解释OSCE","只给总分或主观点评","分值绑定过程证据、关键错误和补练任务"],
        ["武汉大学校本化","通用医学知识库","校本病例、校本量表和人文培养理念"],
        ["教师可控","模型自由生成","病例事实、状态变化和评分规则由教师确认"],
        ["稳定展示","依赖网络和实时模型","离线Demo保障现场演示，正式版可平滑接入模型"],
    ],[1700,3160,4500],8.8)
    add_heading(doc,"建议的竞赛展示主线",2)
    add_para(doc,"用一个患儿呼吸系统病例贯通全部能力：双角色问诊 → 主动查体 → 识别低氧 → 决策改变生命体征 → 与焦虑家长沟通 → 生成OSCE证据报告。评委能够在短时间内理解痛点、看到技术逻辑、体验交互并获得量化结果。")

    add_heading(doc,"八、医学安全、隐私与质量治理",1)
    for t in [
        "用途边界：产品只用于医学教育训练，不直接面向患者提供诊断、治疗或用药建议。",
        "数据边界：真实病例导入前必须完成去标识化、使用授权和最小化处理，不录入姓名、证件号、联系方式、精确住址等信息。",
        "知识边界：指南、共识和教材必须记录版本与来源，设置复核日期；过期内容停止用于正式考核。",
        "评分边界：正式成绩以教师确认的量表和关键错误项为准，模型只承担受控的语义辅助判断。",
        "审计边界：记录病例版本、量表版本、学生行为、系统判断和教师复核，支持争议回看。",
        "生成边界：无法由病例事实支持的信息必须拒绝生成或提示“未提供”，不得补写虚构体征。",
    ]:add_bullet(doc,t)
    add_heading(doc,"质量门槛",2)
    add_table(doc,["质量对象","进入试用前最低要求","进入正式考核前要求"],[
        ["病例","至少1名儿科学教师审核","双人审核、版本锁定、完成试跑"],
        ["评分量表","行为点与分值完整","评分者一致性验证、严重错误项确认"],
        ["知识内容","来源和版本可追溯","到期复核机制生效"],
        ["体征素材","清晰、可辨认、来源合法","授权存档并通过教学适用性审核"],
        ["模型输出","敏感场景有拒答和边界提示","完成红队测试、错误案例复盘和人工抽查"],
    ],[1700,3600,4060],8.8)

    add_heading(doc,"九、实施路径与里程碑",1)
    add_para(doc,"在未获得比赛截止日期前，以T表示教师确认启动日，以D表示比赛最终提交日。拿到比赛通知后，将相对时间转换为具体日期并锁定版本。")
    add_table(doc,["阶段","建议时间","主要工作","阶段交付"],[
        ["P0 竞赛定标","T—T+2工作日","确认赛道、评分细则、参赛名称和展示时长","项目范围、评分映射、材料倒排表"],
        ["P1 内容建模","T+3—T+8工作日","整理病例、量表、沟通场景和知识来源","首批校本病例与评分规则"],
        ["P2 正式实现","D-20—D-12工作日","开发可视化患儿模型、身体部位热点、器材工具栏、匹配校验与反馈动画；接入模型/知识库和审计","可试用版本"],
        ["P3 教师验收","D-11—D-7工作日","医学审核、评分校准、现场演练和修订","参赛候选版本"],
        ["P4 参赛冻结","D-6—D-3工作日","冻结病例与代码，录制视频，准备答辩","最终Demo、视频、PPT和申报书"],
        ["P5 提交备份","最晚D-2工作日","完成平台提交、离线备份和设备验证","可恢复的完整参赛包"],
    ],[1200,1600,3650,2910],8.3)

    add_heading(doc,"十、验收标准",1)
    add_heading(doc,"10.1 汇报文档",2)
    for t in ["老师无需技术背景即可说清项目解决什么问题、如何工作、为何具有儿科和武大特色。","包含总体架构、四库三引擎、完整处理链、Demo路径、治理边界、实施计划和材料清单。","材料清单逐项给出用途、格式、责任建议和最晚提交时间。"]:add_bullet(doc,t)
    add_heading(doc,"10.2 演示Demo",2)
    for t in ["可离线打开，无需API密钥；Chrome/Edge 100%缩放下关键内容完整可见。","至少一个病例完整贯通输入、处理和输出。","教师可以手动操作，也可一键完成标准演示。","OSCE报告显示总分、能力维度、扣分证据和补练建议。","页面明确展示模拟数据、教学用途和非医疗建议边界。"]:add_bullet(doc,t)
    add_heading(doc,"10.3 正式参赛版",2)
    for t in ["病例、量表、指南和体征素材均有教师审核、版本和来源记录。","患儿模型的全部身体部位热点均可点击并触发对应动画或表面体征；不同屏幕尺寸下无明显错位。","听诊器、血压计、体温计等器材必须通过正确部位与操作条件校验后才展示深层信息；错误组合不得泄露结果。","操作反馈覆盖动画、器材状态、成功提示和分级纠错，且全部操作进入OSCE证据记录。","按照部位×器材×病例状态形成测试矩阵，完成正常、错误、重复、越序和边界操作测试。","至少完成5名目标学生试用及1轮教师评分一致性校准。","完成演示设备、网络异常、模型不可用等备份流程。","申报书、PPT、视频与系统中的项目名称、数据和创新点保持一致。"]:add_bullet(doc,t)

    add_heading(doc,"十一、结论与需要老师决策的事项",1)
    add_para(doc,"“珞珈儿科智训”已经具备参赛型作品的基本叙事和可运行演示：以儿科学专业痛点为起点，以双角色、动态病例和证据化OSCE为技术亮点，以校本病例、校本量表和医学人文形成武汉大学特色。下一阶段的关键不在继续增加泛化功能，而在于用教师提供的真实教学材料替换演示数据，并按照比赛评分细则完成定向强化。")
    add_callout(doc,"请老师优先确认", "①比赛通知和截止日期；②正式项目名称；③首批病例与OSCE量表；④负责医学审核的教师；⑤最终落地平台。",fill=PALE_GOLD,accent=GOLD)

    add_heading(doc,"参考依据",1)
    add_para(doc,"1. 武汉大学新闻网：《推动“新医科”建设，培养创新人才》，https://news.whu.edu.cn/info/1751/438867.htm",size=9,color=MUTED)
    add_para(doc,"2. 武汉大学官网：本科生教育与人才培养介绍，https://www.whu.edu.cn/rcpy/bksjy.htm",size=9,color=MUTED)
    add_para(doc,"3. 本项目演示病例、量表和分值为产品功能演示数据，不作为临床或正式教学标准；正式内容以教师审核版本为准。",size=9,color=MUTED)

    doc.add_page_break()
    add_heading(doc,"十二、需要老师配合提供的材料清单（文档末尾）",1)
    add_para(doc,"时间说明：T为教师确认项目启动日，D为大赛最终提交日。若两种时限冲突，以更早者为准。收到大赛正式通知后，应在1个工作日内将本表换算为具体日期。带“必须”的项目未提交前，不进入对应正式内容的冻结或考核使用。")
    materials=[
        ["01 必须","大赛通知、赛道说明、评分细则、申报模板","确定参赛边界、技术重点、申报字段和展示时长","原文件或官方链接","T日；且不晚于D-25工作日"],
        ["02 必须","项目正式名称、参赛成员、指导教师和署名顺序","统一Demo、文档、视频和申报书信息","可复制文字","T+1工作日；且不晚于D-20工作日"],
        ["03 必须","儿科学课程大纲、培养目标、见习/实习要求","建立能力框架并映射四库和OSCE维度","DOCX/PDF，标明适用年级","T+2工作日；且不晚于D-18工作日"],
        ["04 必须","现行儿科OSCE考站任务、评分表和关键错误项","替换Demo评分，保证训练与真实考核一致","可编辑表格优先，附评分说明","T+3工作日；且不晚于D-16工作日"],
        ["05 必须","首批10—20例脱敏教学病例；至少1例完整示范病例","建设校本病例库并制作参赛主病例","使用统一病例模板；不得含身份信息","T+5工作日；且不晚于D-15工作日"],
        ["06 必须","儿童传染病教学重点、指定教材/指南/共识清单","建立年龄差异、隔离、报告和知识版本规则","名称、版本、发布日期、链接或文件","T+5工作日；且不晚于D-14工作日"],
        ["07 必须","典型儿科沟通场景及教师推荐表达","建设家长焦虑、拒检、知情同意、青春期隐私等任务","每场景含背景、目标、禁忌表达、评分点","T+5工作日；且不晚于D-14工作日"],
        ["08 必须","病例和内容审核教师名单及可参加的审核时间","完成医学审核、争议裁决和版本签字","姓名、专长、可用时间，不需私人联系方式","T+2工作日；首次审核不晚于D-10工作日"],
        ["09 必须","身体部位—检查器材—可获取信息—操作顺序映射表","实现器材与部位专业校验，限定表面体征与深层信息的开放条件","可编辑表格；每条规则注明适用年龄、病例和审核人","T+6工作日；且不晚于D-13工作日"],
        ["10 条件必须","患儿模型、身体部位动作、器材图标及皮疹、呼吸音、心音等素材与授权","建设点击动画、器材操作和多模态体征训练；无授权素材不得进入正式系统","原文件＋动作说明＋来源＋使用授权说明","首批T+7工作日；最终不晚于D-10工作日"],
        ["11 必须","武汉大学/学院/附属医院可公开使用的特色介绍和视觉标识授权范围","准确表达校本特色，避免未经授权使用名称、院徽或宣传语","官方文字、VI文件或书面确认","T+3工作日；且不晚于D-15工作日"],
        ["12 必须","真实病例脱敏规范、数据使用范围及伦理/信息部门意见","确定数据能否进入比赛Demo、云端模型或仅限校内环境","现行制度、审批结论或明确负责人意见","任何真实数据导入前；最晚D-14工作日"],
        ["13 必须","最终落地平台与账号条件：扣子/独立网页/校内平台","决定模型、知识库、登录、患儿模型技术路线、部署和网络方案","平台名称、账号权限、网络限制；不得发送密钥","T+2工作日；且不晚于D-18工作日"],
        ["14 建议","学生常见失分点、教师教学痛点、课程评价或访谈摘要","用真实证据说明问题价值并优化补练任务","匿名汇总数据或5—10条教师观察","T+7工作日；且不晚于D-10工作日"],
        ["15 建议","一段教师标准示范或标准问诊与查体脚本","校准角色回答、检查顺序、器材操作、沟通表达和参赛演示台词","文字、录音或视频；需确认使用范围","T+7工作日；且不晚于D-9工作日"],
        ["16 必须","交互式查体功能测试用例及医学验收教师安排","验证全部部位交互、器材逻辑、信息展示和错误提示符合真实检查流程","部位×器材测试矩阵、预期结果、审核人和验收时间","初稿D-12工作日；签字不晚于D-7工作日"],
        ["17 必须","参赛版最终医学内容确认记录","冻结病例、量表、指南和安全声明，防止提交前漂移","签字扫描件或可追溯电子确认","不晚于D-6工作日"],
        ["18 必须","汇报PPT/视频时长、屏幕比例、现场网络与设备条件","确定演示分辨率、离线备份和故障预案","比赛通知截图或现场要求","知悉后立即；最晚D-7工作日"],
    ]
    add_table(doc,["优先级","材料","用途","提交格式要求","最晚提交时间"],materials,[900,2250,2700,1800,1710],7.4)
    add_callout(doc,"材料提交原则", "不要通过非授权渠道传输真实患儿身份信息、病历号、联系方式、影像号、完整住址或任何API密钥。无法确认是否可用的材料，应先提交目录和脱敏样例，由项目组确认后再批量整理。",fill="FFF0F0",accent=RED)

    doc.core_properties.title="珞珈儿科智训：儿童全周期临床胜任力训练与OSCE考核智能体"
    doc.core_properties.subject="智能体大赛教师汇报与项目建设说明"
    doc.core_properties.author="武汉大学儿科学智能体项目组"
    doc.core_properties.keywords="儿科学, 智能体, OSCE, 医学教育, 武汉大学"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
