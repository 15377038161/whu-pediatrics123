from __future__ import annotations

from datetime import date
from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw

import build_report as br


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "实施方案"
ASSETS = OUT / "assets"
OUT.mkdir(parents=True, exist_ok=True)
ASSETS.mkdir(parents=True, exist_ok=True)

DOCX_OUT = OUT / "珞珈儿科智训_智能体实施详细计划与技术方案.docx"
MD_OUT = OUT / "珞珈儿科智训_智能体实施详细计划与技术方案.md"
SELECTED_UI = ROOT / "output" / "ui-renderings" / "04-辅助检查临床判断与处置.png"

NAVY = "0B3558"
BLUE = "176FA9"
TEAL = "148C87"
GOLD = "C47A14"
CORAL = "C64A3B"
INK = "172B3A"
MUTED = "61727D"
LINE = "CDD9E0"
PALE_BLUE = "EAF3F8"
PALE_TEAL = "EAF6F3"
PALE_GOLD = "FFF4DF"
PALE_RED = "FCEDEA"
WHITE = "FFFFFF"


def set_font(run, size=11, color=INK, bold=False, italic=False, name="Microsoft YaHei"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def setup_document(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    specs = {
        1: (16, NAVY, 18, 10),
        2: (13, BLUE, 14, 7),
        3: (12, "245B78", 10, 5),
    }
    for level, (size, color, before, after) in specs.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.add_run("珞珈儿科智训  |  智能体实施详细计划与技术方案"), 8.5, MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("教师确认与研发实施版  ·  "), 8, MUTED)
    br.add_field(footer, "PAGE")


def add_para(doc, text="", *, size=11, color=INK, bold=False, italic=False,
             align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.keep_with_next = keep
    set_font(p.add_run(text), size, color, bold, italic)
    return p


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    set_font(p.add_run(text), 10.7)
    return p


def add_numbered_list(doc, items):
    num_id = br.create_decimal_numbering(doc)
    for text in items:
        p = br.add_number(doc, text, num_id)
        p.paragraph_format.line_spacing = 1.25
        for run in p.runs:
            set_font(run, 10.7)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, bottom=90, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_table(doc, headers, rows, widths, font_size=9.1):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Table Grid"
    set_repeat_header(table.rows[0])
    for i, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_fill(cell, "E8EEF5")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(header), 9.2, NAVY, True)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_idx % 2 == 1:
                set_cell_fill(cell, "F8FAFB")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if i == 0 and len(str(value)) < 16:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(str(value)), font_size, INK, i == 0)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
    br.set_table_geometry(table, [int(x * 1440) for x in widths], indent=120)
    add_para(doc, "", after=2)
    return table


def add_callout(doc, label, text, fill=PALE_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    row_pr = table.rows[0]._tr.get_or_add_trPr()
    row_pr.append(OxmlElement("w:cantSplit"))
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_fill(cell, fill)
    set_cell_margins(cell, 130, 130, 180, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(label + "  "), 10.5, accent, True)
    set_font(p.add_run(text), 10.5, INK)
    br.set_table_geometry(table, [9360], indent=120)
    add_para(doc, "", after=2)


def add_picture(doc, path: Path, caption: str, width=6.45):
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_para(doc, caption, size=8.6, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=7)


def arrow(draw, start, end, color="#6FAFC2", width=6):
    draw.line([start, end], fill=color, width=width)
    x, y = end
    draw.polygon([(x, y), (x - 18, y - 11), (x - 18, y + 11)], fill=color)


def draw_box(draw, xy, title, body="", fill="#FFFFFF", outline="#CDD9E0", title_color="#0B3558"):
    br.rounded(draw, xy, fill, outline, 3, 18)
    x1, y1, x2, y2 = xy
    draw.text((x1 + 24, y1 + 18), title, font=br.font(25, True), fill=title_color)
    if body:
        draw.multiline_text((x1 + 24, y1 + 62), body, font=br.font(18), fill="#61727D", spacing=7)


def build_agent_architecture():
    img = Image.new("RGB", (1800, 1220), "#F5F8FA")
    d = ImageDraw.Draw(img)
    d.text((70, 35), "一个智能体内核，三种教学模式，多端统一呈现", font=br.font(42, True), fill="#0B3558")
    d.text((70, 92), "UI不保存业务真相；所有交互都进入同一智能体会话、病例状态和证据账本。", font=br.font(22), fill="#61727D")

    draw_box(d, (70, 165, 1730, 285), "交互呈现层", "Web学生端  ·  手机学生端  ·  教师工作台  ·  比赛展示大屏", "#FFFFFF")
    arrow(d, (900, 290), (900, 350), "#148C87")
    br.rounded(d, (70, 355, 1730, 695), "#0B3558", "#0B3558", 2, 22)
    d.text((100, 380), "珞珈儿科智训主智能体（Coze Agent）", font=br.font(31, True), fill="#FFFFFF")
    d.text((100, 425), "统一身份、目标、上下文、计划与工具调用；模式变化只改变策略，不更换内核。", font=br.font(20), fill="#CFE6EF")
    agents = [
        ("会话与模式路由", "识别训练/专项/OSCE\n加载相应提示与限制"),
        ("病例状态管理", "阶段、角色、已知信息\n生命体征、解锁条件"),
        ("双角色模拟", "患儿/家长回答边界\n年龄化表达与信息差"),
        ("计划与工具执行", "选择工具、参数校验\n等待结果、继续推理"),
        ("证据评价", "行为事件映射量表\n扣分证据与补练建议"),
    ]
    for i, (title, body) in enumerate(agents):
        x = 105 + i * 320
        draw_box(d, (x, 490, x + 270, 640), title, body, "#FFFFFF", "#9FC6D3")

    arrow(d, (900, 700), (900, 755), "#148C87")
    tools = [
        ("专业工具层", "问诊、查体、器材、辅助检查、诊断、处置、沟通、提交考站"),
        ("知识与规则层", "病例库、体征库、传染病库、沟通库、OSCE量表；教师审核与版本锁定"),
        ("数据与证据层", "会话状态、事件账本、评分证据、学习画像、内容版本、审计日志"),
    ]
    for i, (title, body) in enumerate(tools):
        y = 770 + i * 135
        draw_box(d, (150, y, 1650, y + 105), title, body, ["#EAF3F8", "#EAF6F3", "#FFF4DF"][i])
    path = ASSETS / "智能体总体架构.png"
    img.save(path)
    return path


def build_agent_loop():
    img = Image.new("RGB", (1800, 650), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((70, 35), "智能体每一步都运行同一闭环", font=br.font(40, True), fill="#0B3558")
    steps = [
        ("1 观察", "读取学生输入\n与当前页面动作"),
        ("2 取回", "读取病例状态\n知识与评分规则"),
        ("3 计划", "确定回答者\n决定调用何种工具"),
        ("4 执行", "调用专业工具\n校验参数与顺序"),
        ("5 更新", "推进病例状态\n写入证据事件"),
        ("6 反馈", "返回角色回答\n体征/提示/下一步"),
    ]
    for i, (title, body) in enumerate(steps):
        x = 55 + i * 290
        draw_box(d, (x, 170, x + 245, 380), title, body, "#EAF3F8" if i < 3 else "#EAF6F3")
        if i < len(steps) - 1:
            arrow(d, (x + 248, 275), (x + 282, 275))
    br.rounded(d, (260, 470, 1540, 575), "#FFF4DF", "#E0BD7D", 2, 18)
    d.text((310, 492), "证据账本贯穿全程：谁在什么时间、对什么对象、执行了什么动作、结果是什么、对应哪条量表。", font=br.font(22, True), fill="#7A5311")
    path = ASSETS / "智能体执行闭环.png"
    img.save(path)
    return path


def build_state_machine():
    img = Image.new("RGB", (1800, 900), "#F5F8FA")
    d = ImageDraw.Draw(img)
    d.text((70, 35), "统一病例状态机：训练、专项与OSCE共用", font=br.font(40, True), fill="#0B3558")
    stages = ["接诊信息", "主动问诊", "互动查体", "辅助检查", "临床判断", "处置计划", "家长沟通", "反馈/评分"]
    for i, stage in enumerate(stages):
        x = 55 + (i % 4) * 430
        y = 165 if i < 4 else 520
        fill = "#EAF3F8" if i < 4 else "#EAF6F3"
        draw_box(d, (x, y, x + 345, y + 190), f"{i+1}. {stage}", [
            "加载病例与模式\n创建会话快照", "更新已问项目\n区分患儿/家长", "校验部位、器材\n准备与操作顺序", "记录选择理由\n返回结构化结果",
            "结论必须绑定证据\n不直接泄露答案", "安全规则拦截\n记录隔离/转诊", "评价情绪回应\n通俗解释与确认", "训练即时反馈\nOSCE结束后评分"
        ][i], fill)
        if i in (0,1,2,4,5,6):
            arrow(d, (x + 350, y + 95), (x + 420, y + 95))
    arrow(d, (1620, 360), (1620, 500), "#C47A14")
    d.text((70, 810), "模式策略：训练模式允许提示与重试；专项模式缩小目标与时长；OSCE模式关闭提示、限时且不允许中途重置。", font=br.font(22, True), fill="#7A5311")
    path = ASSETS / "统一病例状态机.png"
    img.save(path)
    return path


def build_deployment():
    img = Image.new("RGB", (1800, 920), "#FFFFFF")
    d = ImageDraw.Draw(img)
    d.text((70, 35), "推荐部署：Coze主智能体 + 统一业务服务 + 响应式前端", font=br.font(40, True), fill="#0B3558")
    cols = [
        (90, "终端", "Web学生端\n手机端/PWA\n教师工作台"),
        (470, "接入与会话", "API网关/BFF\n登录与权限\nSSE流式消息"),
        (850, "智能体运行", "Coze主智能体\n工作流与知识检索\n插件/工具调用"),
        (1230, "业务与数据", "FastAPI模块化单体\nPostgreSQL + Redis\n对象存储与审计"),
    ]
    for x, title, body in cols:
        draw_box(d, (x, 180, x + 300, 450), title, body, "#EAF3F8" if x < 800 else "#EAF6F3")
    for x in (395, 775, 1155):
        arrow(d, (x, 315), (x + 65, 315))
    draw_box(d, (190, 565, 1610, 735), "统一专业工具API", "interview.ask · exam.use_instrument · auxiliary.select · reasoning.submit · treatment.submit · communication.send · station.submit · feedback.generate", "#FFF4DF")
    d.text((190, 790), "首版采用模块化单体，不拆微服务；先保证一个病例闭环、证据可追溯和规则一致，再按负载拆分。", font=br.font(22, True), fill="#7A5311")
    path = ASSETS / "推荐部署架构.png"
    img.save(path)
    return path


def build_roadmap():
    img = Image.new("RGB", (1800, 930), "#F5F8FA")
    d = ImageDraw.Draw(img)
    d.text((70, 35), "16周实施路线：先打通一个完整病例，再扩展能力", font=br.font(40, True), fill="#0B3558")
    phases = [
        ("W1", "范围冻结", "病例、角色、量表\n验收口径"),
        ("W2-3", "底座", "主智能体、状态机\n工具协议、证据账本"),
        ("W4-6", "问诊与查体", "双角色问诊\n部位/器材/操作校验"),
        ("W7-9", "判断与沟通", "辅助检查、诊断证据\n处置与家长沟通"),
        ("W10-11", "OSCE", "限时、留痕\n量表评分与复核"),
        ("W12-13", "专项与教师端", "补弱训练\n资源审核与任务"),
        ("W14-15", "联调与医学验证", "多端、性能、安全\n教师逐例审核"),
        ("W16", "比赛交付", "一例闭环演示\n文档与答辩材料"),
    ]
    for i, (week, title, body) in enumerate(phases):
        x = 55 + (i % 4) * 430
        y = 160 if i < 4 else 520
        fill = "#EAF3F8" if i < 4 else "#EAF6F3"
        draw_box(d, (x, y, x + 345, y + 210), week + "  " + title, body, fill)
        if i in (0,1,2,4,5,6):
            arrow(d, (x + 350, y + 105), (x + 420, y + 105))
    arrow(d, (1620, 375), (1620, 505), "#C47A14")
    d.text((70, 840), "闸门原则：上一阶段必须形成可演示、可测试、可追溯的纵向切片，才进入下一阶段。", font=br.font(22, True), fill="#7A5311")
    path = ASSETS / "实施路线图.png"
    img.save(path)
    return path


def cover(doc):
    add_para(doc, "武汉大学儿科学 · 智能体大赛项目", size=11, color=GOLD, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=42)
    add_para(doc, "珞珈儿科智训", size=30, color=NAVY, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    add_para(doc, "智能体实施详细计划与技术方案", size=17, color=BLUE, bold=True,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
    add_para(doc, "统一智能体内核 · 完整临床闭环 · 可解释OSCE · 教师内容治理", size=11.5, color=TEAL,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=50)
    add_callout(doc, "核心结论", "本项目不建设彼此割裂的问诊、查体、沟通和OSCE系统，而是建设一个能够理解学生意图、维护病例状态、调用专业工具、生成角色反馈并记录评价证据的儿科临床教学智能体。所有页面只是该智能体不同阶段和不同角色的可视化操作面。", PALE_GOLD, GOLD)
    add_para(doc, "方案状态：实施基线确认稿", size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, before=60, after=3)
    add_para(doc, f"编制日期：{date.today().strftime('%Y年%m月%d日')}", size=10, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_page_break()


def add_contents(doc):
    add_heading(doc, "目录与阅读建议", 1)
    sections = [
        "1. 最终架构决策：一个智能体，不是多个功能系统", "2. 建设目标、用户与边界", "3. 智能体总体架构",
        "4. 统一病例状态机与运行模式", "5. 核心智能体能力与专业工具", "6. 知识库、数据模型与证据账本",
        "7. UI与智能体能力的映射", "8. 接口与技术实现", "9. 医学质量、安全与合规",
        "10. 16周实施计划与交付物", "11. 测试、验收与完成定义", "12. 教师材料清单与协同机制",
        "13. 比赛展示方案", "14. 风险、取舍与后续扩展", "附录A. 核心接口清单", "附录B. 关键事件与验收清单"
    ]
    for s in sections:
        add_bullet(doc, s)
    add_callout(doc, "阅读顺序", "老师可优先阅读第1、3、10、12、13节；研发人员重点阅读第4至第11节。", PALE_BLUE, BLUE)
    doc.add_page_break()


def build_docx(diagrams):
    doc = Document()
    setup_document(doc)
    cover(doc)
    add_contents(doc)

    add_heading(doc, "1. 最终架构决策：一个智能体，不是多个功能系统", 1)
    add_callout(doc, "不可变原则", "系统只有一个“珞珈儿科智训主智能体”。模拟病例训练、专项能力训练和OSCE模拟考站是同一智能体的三种运行策略；儿科专科资源库是智能体的内容底座和教师治理面。", PALE_GOLD, GOLD)
    add_heading(doc, "1.1 智能体的判定标准", 2)
    for t in [
        "能够理解学生的自然语言与界面操作，并结合当前病例阶段判断真实意图。",
        "能够维护跨步骤的上下文和病例状态，例如已问病史、已完成查体、已解锁结果、当前角色与剩余时间。",
        "能够自主选择并调用受控的专业工具，而不是仅靠大模型自由生成答案。",
        "能够根据工具结果继续计划下一步，并向学生输出患儿、家长、教学提示或考核反馈。",
        "能够把全过程写入证据账本，使每一条提示、得分和扣分都能回到具体行为。",
        "能够在教师审核的病例、规则和量表边界内工作，不能自行改写正式医学标准。",
    ]: add_bullet(doc, t)
    add_heading(doc, "1.2 明确不采用的做法", 2)
    add_table(doc, ["不采用", "原因", "替代方案"], [
        ("每个页面各写一套业务逻辑", "同一病例在不同页面会出现状态不一致，评分无法追溯", "统一病例状态服务，页面只读写标准动作"),
        ("用一个大提示词包办全部功能", "难以校验器材、顺序、安全规则和量表证据", "主智能体负责计划，确定性工具负责专业动作"),
        ("训练和OSCE复制两套代码", "版本漂移，教师维护双倍成本", "同一内核，使用模式策略开关提示、计时与重试"),
        ("大模型直接给总分", "不可解释、不可复核、竞赛说服力弱", "事件证据先映射量表，再由智能体生成解释"),
        ("学生首页堆放所有资源入口", "稀释主任务，资源库本应服务教师治理", "学生仅保留模拟、专项、OSCE三入口"),
    ], [1.45, 2.35, 2.70], 8.8)

    add_heading(doc, "2. 建设目标、用户与边界", 1)
    add_heading(doc, "2.1 一句话目标", 2)
    add_para(doc, "以标准化儿科病例为载体，让学生在同一个智能体会话中完成问诊、查体、辅助检查、临床判断、处置和家长沟通，并获得有证据、可复核、可补练的OSCE评价。")
    add_heading(doc, "2.2 用户与主要任务", 2)
    add_table(doc, ["角色", "主要任务", "系统承诺"], [
        ("学生", "完成综合病例、专项补弱和OSCE考核", "主动提问、主动操作、按证据学习，不被系统提前泄露答案"),
        ("课程教师", "布置任务、查看班级表现、追加评语", "看到过程而非只看总分，能定位高频漏问和错误操作"),
        ("医学审核教师", "审核病例、素材、标准答案、量表与版本", "正式内容可追溯、可撤回、不可被AI覆盖"),
        ("平台管理员", "账号、权限、资源、审计和运行保障", "最小权限、脱敏、日志与备份"),
    ], [1.15, 2.45, 2.90])
    add_heading(doc, "2.3 首版范围", 2)
    for t in [
        "以1例儿科呼吸/感染综合病例打通完整纵向闭环，作为比赛主展示病例。",
        "实现文字问诊为主、语音输入为辅；患儿和家长双角色回答边界明确。",
        "实现身体部位点击、器材选择、操作准备与器材—部位—顺序校验。",
        "实现结构化辅助检查、诊断证据、鉴别诊断、处置与家长沟通。",
        "实现训练模式即时提示、OSCE模式限时无提示和专项模式短任务补弱。",
        "实现教师资源审核、任务布置、学生过程查看和评分复核。",
    ]: add_bullet(doc, t)
    add_heading(doc, "2.4 首版暂不承诺", 2)
    for t in ["不做真实医疗诊断或面向患者提供诊疗建议。", "不以高度拟真的三维数字人为首版前提，优先保证专业交互与证据闭环。", "不一次覆盖全部儿科病种，先建立可复制的病例制作标准。", "不将AI自动生成内容直接发布为正式教学答案。"]: add_bullet(doc, t)

    add_heading(doc, "3. 智能体总体架构", 1)
    add_picture(doc, diagrams[0], "图1  一个主智能体连接多端、工具、知识与证据")
    doc.add_page_break()
    add_heading(doc, "3.1 四层职责", 2)
    add_table(doc, ["层级", "职责", "关键约束"], [
        ("交互呈现层", "Web、手机、教师端展示状态并收集操作", "不得在前端私自推断检查结果或分数"),
        ("主智能体层", "理解意图、读取上下文、制定计划、调用工具、组织反馈", "所有模式共用同一身份与会话主线"),
        ("专业工具层", "执行问诊、查体、检查、判断、处置、沟通和评分动作", "确定性校验优先于语言模型自由生成"),
        ("知识与数据层", "教师审核内容、病例状态、证据账本、画像和审计", "版本可追溯，正式答案不可被AI覆盖"),
    ], [1.25, 3.05, 2.20])
    add_heading(doc, "3.2 Coze主智能体的职责", 2)
    for t in [
        "保持“儿科临床教学教练/标准化病人编排者”的统一身份。",
        "根据入口与上下文选择训练、专项或OSCE策略，不创建三个互相独立的智能体。",
        "把学生自然语言转换为受控工具参数，例如识别问题主题、回答角色、检查部位和器材。",
        "调用统一业务服务后，根据结构化结果生成自然、年龄适配且不越界的患儿或家长回复。",
        "在训练模式给出分级提示，在OSCE模式严格静默直到提交。",
        "在会话结束时读取证据账本，解释评分并推荐专项训练。",
    ]: add_bullet(doc, t)
    add_heading(doc, "3.3 为什么业务服务仍然必要", 2)
    add_para(doc, "智能体不等于把全部规则写进大模型。器材匹配、操作顺序、感染控制、一票否决、计时、评分计算和版本锁定都需要确定性服务。主智能体负责理解、规划与表达；业务服务负责可验证执行。两者合起来，才是可用于医学教学的智能体产品。")
    add_picture(doc, diagrams[1], "图2  观察—取回—计划—执行—更新—反馈的统一智能体循环")

    add_heading(doc, "4. 统一病例状态机与运行模式", 1)
    add_picture(doc, diagrams[2], "图3  训练、专项与OSCE共用同一病例状态机")
    add_heading(doc, "4.1 会话状态对象", 2)
    add_table(doc, ["状态域", "核心字段", "用途"], [
        ("会话", "session_id、student_id、mode、started_at、deadline", "身份、模式、计时和恢复"),
        ("病例", "case_id、case_version、stage、condition_state", "锁定病例版本与病情阶段"),
        ("角色", "active_role、child_capability、parent_emotion", "决定由患儿还是家长回答"),
        ("病史", "asked_topics、disclosed_facts、omissions", "防止主动泄露与支持完整度评价"),
        ("查体", "region、tool、preparation、sequence、findings", "器材—部位—顺序校验与结果解锁"),
        ("推理", "summary、diagnoses、evidence_links、severity", "验证结论是否有证据支持"),
        ("沟通", "objectives、emotion_response、teach_back", "评价沟通质量与结束条件"),
        ("考核", "rubric_version、events、score_status、review_status", "可解释评分与教师复核"),
    ], [1.05, 3.30, 2.15], 8.7)
    add_heading(doc, "4.2 三种运行策略", 2)
    add_table(doc, ["策略", "提示/重试", "计时与反馈", "同一内核中的差异"], [
        ("模拟病例训练", "可配置提示，可重试", "可选计时，即时反馈", "加载training_policy，允许教学工具返回提示"),
        ("专项能力训练", "聚焦单目标，可重复", "短时任务，立即给出掌握度", "加载skill_policy，只开放相关工具与量表项"),
        ("OSCE模拟考站", "关闭提示，不允许中途重置", "必须计时，结束后反馈", "加载osce_policy，锁定流程、证据和提交规则"),
    ], [1.30, 1.65, 1.65, 1.90], 8.6)

    add_heading(doc, "5. 核心智能体能力与专业工具", 1)
    add_heading(doc, "5.1 主智能体内部能力", 2)
    capabilities = [
        ("意图与模式路由", "识别自然语言问题、界面动作、当前步骤和训练模式，决定回答、提示或工具调用。"),
        ("双角色模拟", "根据年龄与表达能力决定患儿回答；出生史、喂养史、接种史等由家长回答；允许信息差并支持核实。"),
        ("病例动态", "根据已执行操作更新病情阶段、可见体征、生命体征和可用检查结果。"),
        ("临床推理教练", "不直接公布答案，而是检查病情摘要、诊断与证据之间是否一致。"),
        ("沟通教练", "评价情绪回应、通俗解释、尊重患儿、确认理解、行动建议和风险告知。"),
        ("可解释评价", "将事件映射到量表条目，输出分数、遗漏、错误、证据时间点与补练建议。"),
    ]
    for title, body in capabilities:
        add_para(doc, title, bold=True, color=BLUE, after=2)
        add_para(doc, body, after=5)
    add_heading(doc, "5.2 首版专业工具清单", 2)
    add_table(doc, ["工具", "核心输入", "结构化输出", "证据事件"], [
        ("case.start", "病例、版本、模式", "会话快照、首阶段", "SESSION_STARTED"),
        ("interview.ask", "问题文本、目标角色", "回答者、已披露事实、提示", "QUESTION_ASKED"),
        ("exam.observe", "身体部位", "表面可见信息", "REGION_OBSERVED"),
        ("exam.use_instrument", "器材、部位、准备、动作", "校验结果、深层体征", "INSTRUMENT_USED"),
        ("auxiliary.select", "检查项目、选择理由", "适宜性、结果或拒绝原因", "AUX_SELECTED"),
        ("reasoning.submit", "摘要、诊断、证据链接", "支持度、缺证据项", "DIAGNOSIS_SUBMITTED"),
        ("treatment.submit", "处置、隔离、上报、转诊", "安全校验、状态更新", "PLAN_SUBMITTED"),
        ("communication.send", "表达文本、沟通目标", "角色回复、目标完成度", "COMMUNICATION_TURN"),
        ("station.submit", "会话、提交原因", "锁定事件、进入评分", "STATION_SUBMITTED"),
        ("feedback.generate", "事件账本、量表版本", "分项得分、证据、补练", "REPORT_GENERATED"),
    ], [1.45, 1.65, 1.95, 1.45], 8.1)
    add_heading(doc, "5.3 大模型与确定性规则的分工", 2)
    add_table(doc, ["由大模型完成", "由规则/服务完成"], [
        ("理解学生问题、归类问诊主题", "是否允许披露该事实、该事实属于患儿还是家长"),
        ("生成符合年龄与情绪的自然语言", "器材、部位、袖带型号、准备与操作顺序校验"),
        ("对学生表达给出通俗化建议", "安全拦截、一票否决、计时和不可逆提交"),
        ("把结构化证据组织为可读反馈", "事件到量表条目的映射和分数计算"),
        ("在候选资料中进行语义检索", "正式答案、指南版本、发布状态和权限"),
    ], [3.25, 3.25], 9.0)

    add_heading(doc, "6. 知识库、数据模型与证据账本", 1)
    add_heading(doc, "6.1 五类教师治理资源", 2)
    add_table(doc, ["资源", "必须包含", "被哪些能力使用"], [
        ("病例库", "完整病史、体征、检查、标准诊断、处置、沟通重点、量表、来源与版本", "全部训练模式"),
        ("儿童体征库", "年龄、部位、器材、步骤、图像/音频/动画、错误提示、临床意义", "查体与专项训练"),
        ("传染病特点库", "传播、季节、接触史、接种、鉴别、重症预警、隔离与报告", "问诊、判断、处置、OSCE"),
        ("沟通场景库", "角色、情绪、诉求、触发点、策略、禁忌表达、结束条件", "家长沟通与专项训练"),
        ("OSCE量表库", "任务、时间、项目、分值、关键步骤、一票否决、证据要求、复核规则", "OSCE与训练反馈"),
    ], [1.20, 3.20, 2.10], 8.5)
    add_heading(doc, "6.2 证据账本", 2)
    add_para(doc, "证据账本采用仅追加事件模型。每次学生输入、器材操作、检查选择、诊断提交、沟通表达和教师复核都生成不可覆盖的事件。当前页面状态可以从事件重建，评分报告也只引用事件，不引用大模型的隐含判断。")
    add_table(doc, ["字段", "示例", "说明"], [
        ("event_id", "evt_01J...", "全局唯一"), ("session_id", "ses_01J...", "一次训练/考站"),
        ("event_type", "INSTRUMENT_USED", "标准事件类型"), ("actor", "student", "student/agent/teacher/system"),
        ("timestamp", "2026-08-14T09:10:22+08:00", "统一时区"), ("payload", "tool=stethoscope, region=chest", "结构化动作参数"),
        ("result", "valid=true, finding=...", "工具执行结果"), ("source_version", "case:v2.3; rubric:v2.1", "内容与评分版本"),
        ("rubric_links", "PE-03; SAF-01", "对应评分项"), ("trace_id", "trc_...", "跨智能体与服务追踪"),
    ], [1.20, 2.25, 3.05], 8.4)
    add_heading(doc, "6.3 记忆策略", 2)
    for t in [
        "短期记忆：仅保存本次会话所需的病例状态、最近对话和工具结果。",
        "长期学习记忆：只保存学生能力维度、历史错误和训练建议，不保存无关个人敏感信息。",
        "医学标准不作为模型自由记忆，而以版本化资源和规则读取。",
        "OSCE开始后锁定病例与量表版本，结束前不受后续资源修改影响。",
    ]: add_bullet(doc, t)

    add_heading(doc, "7. UI与智能体能力的映射", 1)
    if SELECTED_UI.exists():
        add_picture(doc, SELECTED_UI, "图4  已确认的临床判断与处置UI：页面展示的是同一智能体的当前状态与可执行工具")
    add_table(doc, ["页面", "不是独立系统，而是…", "调用的智能体能力/工具"], [
        ("学生首页", "统一会话入口和任务投影", "case.start、session.resume、recommendation.get"),
        ("病例选择", "向主智能体提交病例与模式参数", "case.search、case.start"),
        ("主动问诊", "双角色模拟的对话操作面", "interview.ask、hint.request"),
        ("互动查体", "查体工具调用与状态反馈面", "exam.observe、exam.use_instrument"),
        ("辅助检查/判断", "证据绑定和临床推理操作面", "auxiliary.select、reasoning.submit"),
        ("处置", "安全规则与计划提交面", "treatment.submit"),
        ("家长沟通", "沟通角色模拟与目标评价面", "communication.send"),
        ("OSCE", "同一病例内核加载严格考核策略", "station.submit、feedback.generate"),
        ("专项训练", "主智能体按薄弱项裁剪后的短任务", "skill.start、feedback.generate"),
        ("教师端", "资源、策略、版本和评分治理面", "resource.review、task.publish、score.review"),
    ], [1.25, 2.65, 2.60], 8.5)
    add_heading(doc, "7.1 多端适配规则", 2)
    for t in [
        "Web端保留多栏证据视图，适合完整病例与教师复盘；手机端每次只突出一个主要任务。",
        "手机端对话只展示最近消息，完整记录使用抽屉；器材使用底部抽屉；证据和量表使用折叠面板。",
        "所有子页面必须提供返回与首页；不可逆操作必须二次确认；下一步按钮只在当前阶段必要条件满足后可用。",
        "UI收到的所有状态均带session_version；版本冲突时重新拉取，不允许前端覆盖较新的病例状态。",
        "轻量动效只表达状态变化：热点呼吸、器材吸附、抽屉滑入、正确/错误反馈；支持减少动态效果。",
    ]: add_bullet(doc, t)

    add_heading(doc, "8. 接口与技术实现", 1)
    add_picture(doc, diagrams[3], "图5  推荐部署架构与专业工具API边界")
    add_heading(doc, "8.1 推荐技术栈", 2)
    add_table(doc, ["范围", "建议", "理由"], [
        ("主智能体", "Coze Agent + 工作流 + 知识库 + 插件", "满足智能体参赛形态，承担会话、规划、检索与工具调用"),
        ("前端", "Vue 3 + TypeScript + Vite + Pinia + PWA", "响应式Web与手机复用组件，适合教学交互"),
        ("业务后端", "Python FastAPI模块化单体", "与模型/规则处理生态一致，首版避免微服务复杂度"),
        ("数据库", "PostgreSQL", "适合版本化资源、会话、事件与评分"),
        ("缓存/会话", "Redis", "计时、幂等、短期状态和限流"),
        ("对象存储", "校内对象存储或兼容S3服务", "体征图像、音频、动画与回放证据"),
        ("实时输出", "SSE优先，必要时WebSocket", "流式回答与状态事件，实施简单"),
        ("监控", "结构化日志 + trace_id + 指标看板", "定位智能体、工具、模型和数据问题"),
    ], [1.15, 2.25, 3.10], 8.5)
    add_heading(doc, "8.2 API设计原则", 2)
    for t in [
        "所有写操作携带session_id、request_id和expected_version，保证幂等与并发安全。",
        "工具API只返回结构化结果和可展示文案槽位，不返回任意HTML。",
        "模型调用不直接写数据库，必须通过工具服务写入并由事件账本确认。",
        "前端不直接访问知识库；主智能体或教师治理服务按权限读取。",
        "每次响应返回next_actions，UI据此决定可用按钮，避免页面自创流程。",
    ]: add_bullet(doc, t)
    add_heading(doc, "8.3 关键响应结构", 2)
    add_callout(doc, "标准响应", "session_id + session_version + stage + visible_state + agent_message + tool_result + next_actions + evidence_refs + safety_flags + trace_id。", PALE_BLUE, BLUE)
    add_heading(doc, "8.4 内容发布链", 2)
    add_numbered_list(doc, [
        "教师新建或导入资源草稿，填写来源、指南版本、授权与适用范围。",
        "系统完成格式校验、缺项检查和AI辅助预审，但不自动发布。",
        "医学审核教师逐项确认病史、体征、标准答案、处置和量表。",
        "审核通过后生成不可变版本号；训练会话引用具体版本。",
        "如发现错误，发布新版本或撤回旧版本；历史会话仍可按原版本复盘。",
    ])

    add_heading(doc, "9. 医学质量、安全与合规", 1)
    add_heading(doc, "9.1 四道防线", 2)
    add_table(doc, ["防线", "控制措施", "失败时处理"], [
        ("内容", "来源标注、教师审核、版本锁定、素材授权", "阻止发布或立即撤回"),
        ("工具", "参数白名单、顺序校验、安全拦截、一票否决", "拒绝执行并记录风险事件"),
        ("模型", "角色边界、禁止泄露答案、引用结构化结果、低温度", "降级为模板反馈或转教师复核"),
        ("运行", "权限、脱敏、审计、备份、限流、异常监控", "隔离会话、告警、恢复与复盘"),
    ], [1.10, 3.35, 2.05], 8.6)
    add_heading(doc, "9.2 必须阻止的高风险行为", 2)
    for t in [
        "未完成手卫生或必要准备却继续侵入性/接触性检查。",
        "器材与部位明显不匹配，或血压袖带型号不适合患儿。",
        "OSCE考核中请求提示、回退或重置已提交步骤。",
        "病例结论缺乏已获得证据，却被系统当作正确答案接受。",
        "AI尝试覆盖教师审核后的标准答案、量表或指南版本。",
        "把模拟训练结果解释为真实患者诊断建议。",
    ]: add_bullet(doc, t)
    add_heading(doc, "9.3 隐私与权限", 2)
    add_para(doc, "病例资料必须去标识化；学生只访问被布置或公开的病例；教师按课程和班级授权；审核、发布、撤回、评分复核均记录操作者与时间。语音、图片和回放素材设置保存期限与删除策略，比赛演示使用虚构或授权素材。")

    add_heading(doc, "10. 16周实施计划与交付物", 1)
    add_picture(doc, diagrams[4], "图6  16周实施路线与阶段闸门")
    add_table(doc, ["阶段", "主要工作", "交付物", "阶段验收"], [
        ("W1 范围冻结", "确定主展示病例、四类角色、量表和素材清单；冻结首版边界", "需求基线、病例蓝图、验收清单", "老师签字确认病例与评分框架"),
        ("W2-3 智能体底座", "建立主智能体、会话、状态机、工具协议、事件账本和权限", "可创建/恢复会话的纵向切片", "同一session跨页面状态一致，事件可重放"),
        ("W4-6 问诊与查体", "双角色问诊、年龄表达、部位点击、器材与准备校验", "问诊/查体完整流程", "不问不泄露；正确器材解锁；错误有证据"),
        ("W7-9 判断与沟通", "辅助检查、证据绑定、诊断、处置、家长沟通", "综合病例训练闭环", "结论可追溯到证据，沟通有目标评价"),
        ("W10-11 OSCE", "考站策略、计时、无提示、量表映射、报告与复核", "OSCE考站和可解释报告", "每个扣分有事件、时间和量表条目"),
        ("W12-13 专项与教师端", "补弱任务、能力画像、任务布置、资源审核与班级分析", "专项训练和教师工作台", "OSCE失分可一键进入对应专项"),
        ("W14-15 验证", "功能、医学、浏览器、手机、性能、安全、可访问性验证", "缺陷清单、测试报告、医学审核记录", "P0/P1为0，主病例100%通过"),
        ("W16 比赛交付", "固化演示数据、答辩流程、容灾与离线预案", "演示包、汇报文档、操作手册", "15分钟内完整演示，断网有预案"),
    ], [1.05, 2.25, 1.75, 1.45], 7.9)
    add_heading(doc, "10.1 首个纵向切片的完成顺序", 2)
    add_numbered_list(doc, [
        "学生从首页选择固定病例，创建唯一session。",
        "学生输入一个问题，主智能体判断由患儿或家长回答并写入证据。",
        "学生选择胸部和听诊器，规则服务校验并返回呼吸音结果。",
        "学生提交诊断与依据，系统检查是否引用已获得证据。",
        "学生完成一轮家长沟通，系统记录沟通目标完成情况。",
        "提交后按最小量表生成可解释报告，并跳转到一个专项补练。",
    ])
    add_callout(doc, "实施纪律", "先完成这条端到端链路，再扩充病例、素材、动画和数据分析；任何只完成单页或单模块的进度都不能称为智能体闭环完成。", PALE_RED, CORAL)

    add_heading(doc, "11. 测试、验收与完成定义", 1)
    add_heading(doc, "11.1 测试分层", 2)
    add_table(doc, ["测试层", "重点", "代表用例"], [
        ("单元测试", "状态转移、规则、评分、权限", "错误器材不解锁；量表分数可重算"),
        ("契约测试", "智能体与工具API的输入输出", "字段缺失、重复request_id、版本冲突"),
        ("对话评测", "角色、边界、年龄表达、不泄露", "低龄儿童问题转家长；未问信息不主动回答"),
        ("病例回放", "同一事件序列结果确定", "重放后状态和分数与原会话一致"),
        ("端到端测试", "Web/手机完整业务路径", "病例选择至反馈和专项推荐"),
        ("医学验收", "内容正确、操作规范、量表合理", "教师逐步核对主展示病例"),
        ("非功能测试", "性能、安全、恢复、兼容、可访问性", "并发、断线续训、权限越界、键盘与触控"),
    ], [1.20, 2.30, 3.00], 8.5)
    add_heading(doc, "11.2 关键验收指标", 2)
    for t in [
        "主展示病例从进入到报告的成功率为100%，不存在无法继续的死状态。",
        "问诊事实泄露率为0：未询问的关键病史不由智能体主动披露。",
        "器材—部位—准备—顺序规则覆盖首例全部关键查体动作，错误动作不返回深层结果。",
        "OSCE每个得分/扣分均可定位到事件、时间、量表版本和原始学生行为。",
        "Web与手机端的核心路径均可完成；手机端不出现横向滚动和不可触达按钮。",
        "教师审核后的标准答案不可被AI或普通教师直接覆盖。",
        "关键接口P95响应目标：确定性工具不高于800ms；流式首字不高于3s（模型服务正常时）。",
        "P0/P1缺陷为0；医学审核意见全部关闭或书面接受风险。",
    ]: add_bullet(doc, t)
    add_heading(doc, "11.3 Definition of Done", 2)
    add_para(doc, "一个功能只有同时满足以下条件才算完成：代码合并、接口文档更新、规则与内容版本明确、自动测试通过、真实浏览器/手机流程通过、事件证据可查询、日志无敏感信息、教师可按验收脚本复现、旧入口和重复逻辑已清理。")

    add_heading(doc, "12. 教师材料清单与协同机制", 1)
    add_table(doc, ["材料", "用途", "建议提交时间", "最低要求"], [
        ("主展示病例完整病历", "建立首例状态机与标准答案", "W1第2个工作日前", "去标识化，含完整病史、查体、检查、诊断、处置"),
        ("患儿/家长角色设定", "双角色回答边界与情绪状态", "W1第3个工作日前", "年龄、表达能力、家长身份、信息差、禁答项"),
        ("儿科问诊要点", "建立必问、可选、漏问和诱导性提问规则", "W1结束前", "标注由患儿或家长回答及评分权重"),
        ("体格检查标准", "部位、器材、准备和顺序校验", "W2第2个工作日前", "正常/异常结果、错误提示和安全项"),
        ("体征图片/音频/动画", "可视化与听诊结果", "W3结束前", "来源、授权、适用年龄与病例版本"),
        ("辅助检查与临床路径", "检查适宜性、结果、诊断证据与处置", "W3结束前", "标准答案、可接受替代及不合理选择"),
        ("沟通场景脚本", "家长情绪、诉求、触发点和结束条件", "W5结束前", "推荐表达、禁忌表达、必须告知信息"),
        ("OSCE量表", "事件到评分条目的映射", "W6结束前", "分值、关键步骤、一票否决、证据要求"),
        ("指南/教材依据", "知识库来源与版本治理", "随相应材料同时", "名称、版本/年份、章节或页码"),
        ("班级与任务样例", "教师端和数据分析测试", "W10结束前", "虚构或脱敏学生、班级、截止时间"),
        ("审核教师名单", "资源与评分复核流程", "W1结束前", "课程教师、医学审核教师、最终签字人"),
        ("比赛规则和答辩限制", "演示时长、网络、设备和评审重点", "拿到后24小时内", "官方文件或截图，不依赖口头转述"),
    ], [1.35, 2.05, 1.35, 1.75], 7.8)
    add_heading(doc, "12.1 教师评审节奏", 2)
    for t in [
        "每周一次30分钟内容评审：只确认病例、规则、量表和医学风险。",
        "每两周一次可操作版本验收：老师按脚本亲自完成关键路径。",
        "所有意见进入同一问题清单，标注严重度、责任人、目标版本和关闭证据。",
        "医学分歧由指定最终审核教师裁决，研发不自行选择答案。",
    ]: add_bullet(doc, t)

    add_heading(doc, "13. 比赛展示方案", 1)
    add_heading(doc, "13.1 12—15分钟主展示链", 2)
    demo_steps = [
        ("1分钟", "一句话定位", "不是问答机器人，而是能维护病例状态、调用专业工具并给出证据评分的儿科教学智能体。"),
        ("3分钟", "双角色问诊", "学生主动询问；智能体判断患儿/家长回答；演示信息不主动泄露。"),
        ("3分钟", "可视化查体", "选择部位与器材；错误操作被校验；正确动作解锁深层结果。"),
        ("2分钟", "判断与处置", "诊断必须绑定已获得证据；演示隔离/转诊等规则。"),
        ("2分钟", "家长沟通", "智能体模拟焦虑家长并评价通俗表达和风险告知。"),
        ("2分钟", "OSCE报告", "展示事件、时间、量表、扣分证据与教师复核。"),
        ("1分钟", "闭环", "由失分项自动进入专项补弱，教师端查看班级高频问题。"),
    ]
    add_table(doc, ["时间", "展示节点", "评委应看到的价值"], demo_steps, [0.90, 1.55, 4.05], 8.7)
    add_heading(doc, "13.2 必须准备的容灾", 2)
    for t in ["固定演示账号与固定病例版本。", "模型不可用时切换到已审核的规则化角色回复，仍保留工具与证据闭环。", "准备本地录屏，但现场优先真实操作。", "预置重置会话按钮，仅供演示管理员使用，不出现在OSCE学生界面。", "所有素材本地或校内缓存，避免临时外链失效。"]: add_bullet(doc, t)

    add_heading(doc, "14. 风险、取舍与后续扩展", 1)
    add_table(doc, ["风险", "早期信号", "处置"], [
        ("做成页面集合", "模块各自维护状态或评分", "强制所有操作走工具API和事件账本；禁止前端写结论"),
        ("AI回答越界", "主动给出未问病史或答案", "事实白名单、角色边界、对话评测与模板降级"),
        ("医学内容不足", "开发等待素材或只能用通用答案", "W1锁定首例，材料分批但明确截止时间"),
        ("评分不可解释", "报告只有总分和泛化建议", "量表先定义证据要求，再实现评分"),
        ("过度追求动画", "视觉投入高但专业规则未完成", "首版写实静态+关键动作，动画不阻塞闭环"),
        ("平台锁定", "业务规则只能在提示词中维护", "工具与数据协议独立，Coze通过插件调用"),
        ("性能波动", "模型慢导致页面卡住", "SSE、超时、重试、缓存与规则化降级"),
        ("内容版权/隐私", "素材无授权或病例可识别", "授权台账、去标识化、发布前检查"),
    ], [1.25, 2.20, 3.05], 8.5)
    add_heading(doc, "14.1 后续扩展顺序", 2)
    add_numbered_list(doc, [
        "按课程优先级复制病例模板，逐步扩展呼吸、消化、感染、神经、血液和免疫病例。",
        "增加听诊音频定位、口腔局部放大、姿势变化等体征交互，但仍服从统一工具协议。",
        "增加教师共建、双人审核、版本对比和内容定期复核提醒。",
        "增加更细的能力画像和班级教学诊断，不改变事件账本口径。",
        "必要时再拆分高负载服务；首版保持模块化单体，避免分布式复杂度。",
    ])

    add_heading(doc, "附录A. 核心接口清单", 1)
    add_table(doc, ["方法与路径", "用途", "关键字段", "返回"], [
        ("POST /agent/sessions", "创建训练会话", "case_id, mode, case_version", "session, state, next_actions"),
        ("GET /agent/sessions/{id}", "恢复会话", "session_id", "visible_state, version"),
        ("POST /agent/sessions/{id}/messages", "学生输入问诊/沟通", "text, target_role?, request_id", "stream, role, evidence_refs"),
        ("POST /agent/sessions/{id}/actions", "执行界面动作", "action_type, params, expected_version", "tool_result, state, next_actions"),
        ("POST /agent/sessions/{id}/submit", "提交训练/考站", "reason, expected_version", "locked_state, report_job"),
        ("GET /reports/{id}", "查询反馈报告", "report_id", "scores, evidence, recommendations"),
        ("GET /cases", "检索可用病例", "age, disease, difficulty, focus", "versioned case summaries"),
        ("POST /teacher/tasks", "布置训练", "course, class, case, policy, deadline", "task"),
        ("POST /teacher/resources/{id}/review", "审核资源", "decision, comment, version", "review record"),
        ("POST /teacher/scores/{id}/review", "复核评分", "item, change, reason", "reviewed score + audit"),
    ], [1.95, 1.20, 2.10, 1.25], 7.7)

    add_heading(doc, "附录B. 关键事件与最终验收清单", 1)
    add_heading(doc, "B.1 最小事件集合", 2)
    events = ["SESSION_STARTED", "QUESTION_ASKED", "FACT_DISCLOSED", "HINT_REQUESTED", "REGION_OBSERVED", "INSTRUMENT_USED", "SAFETY_BLOCKED", "AUX_SELECTED", "DIAGNOSIS_SUBMITTED", "PLAN_SUBMITTED", "COMMUNICATION_TURN", "STATION_SUBMITTED", "RUBRIC_ITEM_SCORED", "REPORT_GENERATED", "TEACHER_REVIEWED"]
    add_para(doc, " · ".join(events), size=9.4, color=BLUE)
    add_heading(doc, "B.2 最终验收清单", 2)
    checks = [
        "一个主智能体身份贯穿全部学生路径。", "训练/专项/OSCE共用同一病例状态与工具协议。",
        "学生主动提问，患儿/家长只回答已询问内容。", "部位、器材、准备和顺序校验有效。",
        "诊断与处置绑定已获得证据。", "OSCE关闭提示且结束后统一反馈。",
        "每个得分/扣分可追溯到事件与量表版本。", "失分项能跳转到对应专项训练。",
        "教师可审核、发布、撤回和查看版本。", "Web与手机端均能完成核心路径。",
        "医学内容已脱敏、授权、审核并锁定。", "比赛演示和模型故障降级预案可用。",
    ]
    for t in checks:
        p = add_bullet(doc, "□ " + t)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
    add_callout(doc, "最终判断", "只有当同一病例能够在同一session中完成“问诊—查体—判断—处置—沟通—评分—补练”，且全过程证据可重放，项目才可以称为儿科临床教学智能体已完成。", PALE_GOLD, GOLD)

    doc.save(DOCX_OUT)


def build_markdown():
    text = dedent(f"""
    # 珞珈儿科智训——智能体实施详细计划与技术方案

    > 编制日期：{date.today().strftime('%Y年%m月%d日')}
    >
    > 方案状态：实施基线确认稿

    ## 结论

    本项目建设的是**一个儿科临床教学主智能体**，不是问诊、查体、沟通和OSCE的页面集合。模拟病例训练、专项能力训练和OSCE模拟考站是同一智能体的三种策略；儿科专科资源库是内容底座和教师治理面。所有UI只负责展示智能体状态、收集学生动作，不得自行生成检查结果、业务状态或评分。

    ## 一、统一智能体架构

    - 主智能体：统一身份、会话、意图理解、计划、知识检索、工具调用和自然语言反馈。
    - 统一病例状态机：接诊信息 → 主动问诊 → 互动查体 → 辅助检查 → 判断 → 处置 → 家长沟通 → 反馈/评分。
    - 专业工具：`case.start`、`interview.ask`、`exam.observe`、`exam.use_instrument`、`auxiliary.select`、`reasoning.submit`、`treatment.submit`、`communication.send`、`station.submit`、`feedback.generate`。
    - 知识与规则：病例库、儿童体征库、儿科传染病特点库、沟通场景库、OSCE量表库。
    - 数据与证据：会话快照、仅追加事件账本、评分证据、能力画像、内容版本和审计日志。

    智能体每一步运行“观察—取回—计划—执行—更新—反馈”闭环。大模型负责理解、规划和表达；确定性服务负责器材、部位、顺序、安全、计时、评分和版本锁定。

    ## 二、三种模式共用一个内核

    | 模式 | 提示与重试 | 反馈 | 策略差异 |
    |---|---|---|---|
    | 模拟病例训练 | 可配置提示，可重试 | 即时反馈 | `training_policy` |
    | 专项能力训练 | 聚焦单目标，可重复 | 即时掌握度 | `skill_policy` |
    | OSCE模拟考站 | 无提示，不可重置 | 结束后反馈 | `osce_policy` |

    ## 三、首版实施范围

    1. 以一例儿科呼吸/感染综合病例贯通完整闭环。
    2. 学生主动文字问诊，智能体自动判断患儿或家长回答。
    3. 点击患儿部位，选择器材，校验手卫生、部位、器材和操作顺序。
    4. 完成辅助检查、病情摘要、诊断依据、鉴别诊断、危重程度和处置。
    5. 完成家长沟通目标与表达评价。
    6. 同一病例切换OSCE策略，生成有事件证据的报告并推荐专项补弱。
    7. 教师端完成资源审核、版本发布、任务布置、过程回放和评分复核。

    ## 四、技术方案

    - 主智能体：Coze Agent + 工作流 + 知识库 + 插件。
    - 前端：Vue 3 + TypeScript + Vite + Pinia + PWA。
    - 后端：Python FastAPI模块化单体；首版不拆微服务。
    - 数据：PostgreSQL保存资源、版本、会话、事件和评分；Redis保存计时、幂等和短期状态；对象存储保存图像、音频、动画与回放证据。
    - 实时交互：SSE优先，必要时使用WebSocket。
    - 关键原则：所有写操作必须带`session_id`、`request_id`和`expected_version`；模型不能直接写库；UI只能根据`next_actions`开放操作。

    ## 五、16周实施计划

    | 周期 | 阶段 | 交付与闸门 |
    |---|---|---|
    | W1 | 范围冻结 | 主病例、角色、量表、材料与验收基线完成教师确认 |
    | W2-3 | 智能体底座 | 主智能体、状态机、工具协议、事件账本；状态可重放 |
    | W4-6 | 问诊与查体 | 双角色、不泄露、部位器材顺序校验 |
    | W7-9 | 判断与沟通 | 诊断证据、处置安全、家长沟通闭环 |
    | W10-11 | OSCE | 限时无提示、事件评分、可解释报告与复核 |
    | W12-13 | 专项与教师端 | 自动补弱、任务布置、资源审核与班级分析 |
    | W14-15 | 联调验证 | 医学、功能、多端、性能、安全和可访问性验收 |
    | W16 | 比赛交付 | 固定演示病例、答辩材料、容灾与离线预案 |

    ## 六、完成定义

    只有当同一病例能够在同一`session`中完成“问诊—查体—判断—处置—沟通—评分—补练”，每次动作均写入证据账本，报告能够回到事件、时间和量表版本，项目才可以称为儿科临床教学智能体完成。单页、单模块或只能对话的版本均不算完成。

    ## 七、教师需优先提供的材料

    - W1：主展示病例、患儿/家长角色、问诊要点、审核教师名单。
    - W2-W3：体格检查标准、体征素材及授权、辅助检查和临床路径。
    - W5-W6：沟通场景、OSCE量表、评分证据要求。
    - 全程：对应教材/指南名称、版本、章节或页码；所有病例去标识化。

    详细的模块、接口、数据结构、测试、验收、教师材料用途与比赛演示脚本，以同目录DOCX/PDF正式版为准。
    """).strip() + "\n"
    MD_OUT.write_text(text, encoding="utf-8")


def main():
    diagrams = [build_agent_architecture(), build_agent_loop(), build_state_machine(), build_deployment(), build_roadmap()]
    build_docx(diagrams)
    build_markdown()
    print(DOCX_OUT)
    print(MD_OUT)


if __name__ == "__main__":
    main()
