"""Synthetic regression tests; no private patient sources required."""
import importlib.util
from pathlib import Path
import unittest
from docx import Document

spec = importlib.util.spec_from_file_location('intake', Path(__file__).with_name('import-standardized-case-docx.py'))
intake = importlib.util.module_from_spec(spec)
spec.loader.exec_module(intake)


class IntakeTests(unittest.TestCase):
    def setUp(self):
        self.sha = next(iter(intake.EXPECTED))
        cid, age, sex = intake.EXPECTED[self.sha]
        self.record = dict(caseId=cid, age=age, sex=sex, complaint='测试主诉',
                           documents=[], sections=dict(history='原始病史'))
        self.fields = {'性别、年龄': '男，12岁', '主诉': '测试主诉', '出生史': '测试出生史'}

    def test_completed_right_column_wins_over_old_example(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '性别、年龄：女，1岁6月'
        table.cell(0, 1).text = '男，12岁'
        table.cell(1, 0).text = '主诉：其他患儿模板'
        table.cell(1, 1).text = '测试主诉'
        self.assertEqual(intake.selected_basic_facts(doc)['性别、年龄'], '男，12岁')
        self.assertEqual(intake.selected_basic_facts(doc)['主诉'], '测试主诉')

    def test_rejects_different_patient_or_complaint(self):
        for field, value in [('性别、年龄', '女，1岁6月'), ('主诉', '不同主诉')]:
            with self.subTest(field=field), self.assertRaises(ValueError):
                intake.candidate(self.record, {**self.fields, field: value}, self.sha)

    def test_selected_facts_only_and_idempotent(self):
        fields = {**self.fields, '诊断': '不可照搬', '角色': '其他年龄台词', '既往史': '未经审核'}
        result, additions = intake.candidate(self.record, fields, self.sha)
        self.assertEqual(additions, ['出生史'])
        self.assertEqual(self.record['sections']['history'], '原始病史')
        self.assertNotIn('不可照搬', result['documents'][0]['text'])
        self.assertNotIn('其他年龄', result['documents'][0]['text'])
        self.assertEqual(intake.candidate(result, fields, self.sha), (result, []))


if __name__ == '__main__':
    unittest.main()
