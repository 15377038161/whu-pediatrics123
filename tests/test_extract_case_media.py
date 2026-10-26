"""Source extraction fixtures only; no external services or real patient data."""
import base64
import importlib.util
import io
from hashlib import sha256
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import NameObject, NumberObject, DictionaryObject, DecodedStreamObject

SCRIPTS = Path(__file__).resolve().parents[1] / 'scripts'
spec = importlib.util.spec_from_file_location('extract_media_target', SCRIPTS / 'extract-case-drafts.py')
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)
PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jBfkAAAAASUVORK5CYII=')

class EmbeddedMediaExtractionTest(unittest.TestCase):
    def test_docx_preserves_source_and_keeps_media_private(self):
        with tempfile.TemporaryDirectory(prefix='source-media-docx-') as name:
            root = Path(name); source = root / 'source.docx'
            doc = Document(); doc.add_paragraph('合成资料'); doc.add_picture(io.BytesIO(PNG)); doc.save(source)
            original = source.read_bytes()
            with patch.object(module, 'OUTPUT', root / 'drafts'):
                result = module.extract_file(source)
            self.assertIn('合成资料', result['text'])
            self.assertEqual(result['mediaExtractionStatus'], 'extracted')
            self.assertEqual(len(result['embeddedMediaCandidates']), 1)
            item = result['embeddedMediaCandidates'][0]
            self.assertEqual(Path(item['localPath']).read_bytes(), PNG)
            self.assertFalse(item['bindingAllowed'])
            self.assertEqual(item['privacyReview'], 'not_verified')
            self.assertEqual(item['sourceDocumentSha256'], sha256(original).hexdigest())
            self.assertEqual(source.read_bytes(), original)

    def test_image_only_pdf_is_not_silently_lost(self):
        with tempfile.TemporaryDirectory(prefix='source-media-pdf-') as name:
            root = Path(name); source = root / 'source.pdf'; writer = PdfWriter()
            for _ in range(2):
                page = writer.add_blank_page(width=10, height=10)
                image = DecodedStreamObject(); image.set_data(bytes([255, 0, 0]))
                image.update({NameObject('/Type'): NameObject('/XObject'), NameObject('/Subtype'): NameObject('/Image'), NameObject('/Width'): NumberObject(1), NameObject('/Height'): NumberObject(1), NameObject('/ColorSpace'): NameObject('/DeviceRGB'), NameObject('/BitsPerComponent'): NumberObject(8)})
                page[NameObject('/Resources')] = DictionaryObject({NameObject('/XObject'): DictionaryObject({NameObject('/Im0'): writer._add_object(image)})})
            with source.open('wb') as handle: writer.write(handle)
            original = source.read_bytes()
            with patch.object(module, 'OUTPUT', root / 'drafts'):
                result = module.extract_file(source)
            self.assertEqual(result['status'], 'no-text-ocr-needed')
            self.assertEqual(len(result['embeddedMediaCandidates']), 2)
            self.assertEqual(len({i['sha256'] for i in result['embeddedMediaCandidates']}), 1)
            self.assertEqual(len({i['sourceLocator'] for i in result['embeddedMediaCandidates']}), 2)
            self.assertTrue(all(not i['bindingAllowed'] for i in result['embeddedMediaCandidates']))
            self.assertEqual(source.read_bytes(), original)


    def test_corrupt_private_output_is_reported_without_overwriting(self):
        with tempfile.TemporaryDirectory(prefix='source-media-corrupt-') as name:
            root = Path(name); source = root / 'source.docx'
            doc = Document(); doc.add_paragraph('合成资料'); doc.add_picture(io.BytesIO(PNG)); doc.save(source)
            with patch.object(module, 'OUTPUT', root / 'drafts'):
                first = module.extract_file(source)
                target = Path(first['embeddedMediaCandidates'][0]['localPath'])
                target.write_bytes(b'synthetic-corrupted-output')
                second = module.extract_file(source)
            self.assertEqual(second['mediaExtractionStatus'], 'extract-error')
            self.assertEqual(second['mediaErrorType'], 'ValueError')
            self.assertEqual(target.read_bytes(), b'synthetic-corrupted-output')
            self.assertIn('合成资料', second['text'])

if __name__ == '__main__': unittest.main()
